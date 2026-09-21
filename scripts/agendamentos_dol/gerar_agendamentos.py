"""Gera a planilha de agendamentos Cível/Família da 9ª DPE Ribeirão Preto (janela de 7 dias).

Uso: python gerar_agendamentos.py [caminho_saida.xlsx]
Lê dados/agendamentos.json; compara com a conferência mais recente de dia anterior (dados/conferencias/);
preserva a coluna "Observações da equipe" da planilha existente; registra a execução em dados/historico.json.
"""
import json
import shutil
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = Path(__file__).parent
DADOS = BASE / "dados"
CONFERENCIAS = DADOS / "conferencias"
PRAZOS = BASE.parent / "prazos_dol" / "dados" / "providencias.json"
SAIDA_PADRAO = Path.home() / "Downloads" / "Agendamentos 9ª DPE" / "Agendamentos 9ª DPE.xlsx"

VERDE = "009632"
COR_EDITAVEL = "FFFDE7"
GRUPOS = ["FFFFFF", "EEF6F1"]
SITUACOES = {
    "Novo desde a última conferência": ("C6EFCE", "006100"),
    "Horário alterado": ("FFE699", "7F6000"),
    "Sem mudança": (None, "595959"),
    "Primeira conferência": ("DDEBF7", "1F4E79"),
    "Não aparece mais": ("F2F2F2", "808080"),
}
FINO = Side(style="thin", color="D9D9D9")
BORDA = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)
ABAS = {"presencial": "Presencial", "virtual": "Virtual Cível-Família"}
CAB = [("Data / hora", 12), ("Assistido / repr. legal", 24), ("Nº DOL", 10), ("Processo (CNJ)", 25), ("PA", 13),
       ("Vara", 20), ("Motivo do agendamento", 42), ("Resumo do processo", 55), ("Últimos atendimentos", 55),
       ("Alertas", 34), ("Pendências no DOL", 30), ("Situação na conferência", 18), ("Observações da equipe", 40),
       ("ID", 4)]
COL_OBS, COL_ID = 13, 14


def parse_dt(s):
    return datetime.strptime(s, "%d/%m/%Y %H:%M")


def celula(ws, r, c, v, bold=False, cor=None, fundo=None, h="left", size=10, wrap=True, borda=True, strike=False):
    x = ws.cell(row=r, column=c, value=v)
    x.font = Font(bold=bold, color=cor, size=size, name="Calibri", strike=strike)
    if fundo:
        x.fill = PatternFill("solid", fgColor=fundo)
    x.alignment = Alignment(horizontal=h, vertical="top", wrap_text=wrap)
    if borda:
        x.border = BORDA
    return x


def faixa(ws, titulo, subtitulo, ultima_col):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ultima_col)
    x = celula(ws, 1, 1, titulo, bold=True, cor="FFFFFF", fundo=VERDE, size=15, borda=False)
    x.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 32
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ultima_col)
    celula(ws, 2, 1, subtitulo, cor="595959", size=10, borda=False, wrap=False)
    ws.sheet_view.showGridLines = False


def conferencia_anterior(hoje):
    if not CONFERENCIAS.exists():
        return None
    antigas = sorted(f for f in CONFERENCIAS.glob("*.json") if f.stem < hoje.isoformat())
    return json.loads(antigas[-1].read_text(encoding="utf-8")) if antigas else None


def classificar(atual, anterior, hoje):
    """Marca a situação de cada agendamento e devolve os que sumiram da janela (ainda no futuro)."""
    if anterior is None:
        for a in atual:
            a["situacao"] = "Primeira conferência"
        return []
    antes = {a["agid"]: a for a in anterior["agendamentos"]}
    ids_atuais = {a["agid"] for a in atual}
    for a in atual:
        velho = antes.get(a["agid"])
        if not velho:
            a["situacao"] = "Novo desde a última conferência"
        elif velho["data"] != a["data"]:
            a["situacao"] = "Horário alterado"
            a["situacao_obs"] = f"antes: {velho['data']}"
        else:
            a["situacao"] = "Sem mudança"
    sumidos = []
    for agid, velho in antes.items():
        if agid not in ids_atuais and parse_dt(velho["data"]).date() >= hoje:
            sumidos.append({**velho, "situacao": "Não aparece mais"})
    return sumidos


def carregar_observacoes(pasta):
    notas = {}
    caminho = DADOS / "observacoes.json"
    if caminho.exists():
        notas = json.loads(caminho.read_text(encoding="utf-8"))
    for arq in sorted((f for f in pasta.glob("*.xlsx") if not f.name.startswith("~$")), key=lambda f: f.stat().st_mtime):
        try:
            wb = load_workbook(arq, read_only=True, data_only=True)
        except Exception:
            continue
        for aba in ABAS.values():
            if aba not in wb.sheetnames:
                continue
            for row in wb[aba].iter_rows(min_row=5, values_only=True):
                if len(row) >= COL_ID and row[COL_ID - 1]:
                    notas[row[COL_ID - 1]] = (row[COL_OBS - 1] or "").strip() if isinstance(row[COL_OBS - 1], str) else ""
    return {k: v for k, v in notas.items() if v}


def pendencias_dol():
    if not PRAZOS.exists():
        return {}
    provs = json.loads(PRAZOS.read_text(encoding="utf-8")).get("providencias", [])
    mapa = {}
    for p in provs:
        texto = p.get("acao") or p.get("pedido", "")[:100]
        prazo = p.get("prazo_providencia") or "sem data"
        mapa.setdefault(p["processo"], []).append(f"{texto} (prazo da providência: {prazo})")
    return mapa


def aba_agendamentos(wb, modo, agendamentos, notas, pend, periodo):
    ws = wb.create_sheet(ABAS[modo])
    lista = sorted([a for a in agendamentos if a["modo"] == modo], key=lambda a: (parse_dt(a["data"]), a["nome"]))
    n_ativos = sum(1 for a in lista if a["situacao"] != "Não aparece mais")
    faixa(ws, f"Agendamentos {ABAS[modo].replace('-', '/')}  |  9ª DPE Ribeirão Preto",
          f"Janela de {periodo['inicio']} a {periodo['fim']}   |   {n_ativos} pessoas da 9ª agendadas   |   "
          f"Atualizado em {datetime.now().strftime('%d/%m/%Y às %H:%M')}   |   A coluna amarela é da equipe", len(CAB) - 1)
    for i, (h, w) in enumerate(CAB, 1):
        cor_cab = "BF9000" if i == COL_OBS else VERDE
        celula(ws, 4, i, h, bold=True, cor="FFFFFF", fundo=cor_cab, h="center", size=10)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.column_dimensions[get_column_letter(COL_ID)].hidden = True
    ws.row_dimensions[4].height = 30

    r = 5
    for g, a in enumerate(lista):
        fundo = GRUPOS[g % 2]
        sumiu = a["situacao"] == "Não aparece mais"
        bg_sit, fg_sit = SITUACOES[a["situacao"]]
        nome = a["nome"] + (f"\n({a['obs_nome']})" if a.get("obs_nome") else "")
        for j, p in enumerate(a["processos"]):
            primeira = j == 0
            chave = f"{a['agid']}|{p['autos']}"
            situacao = a["situacao"] + (f" ({a['situacao_obs']})" if a.get("situacao_obs") else "")
            valores = [a["data"] if primeira else "", nome if primeira else "", a["dol"] if primeira else "",
                       p["autos"], p["pa"], p.get("vara", ""), a.get("motivo", "") if primeira else "",
                       p.get("resumo", ""), p.get("atendimentos", ""), p.get("alerta", ""),
                       "\n".join(pend.get(p["autos"], [])), situacao if primeira else "", notas.get(chave, ""), chave]
            for c, v in enumerate(valores, 1):
                x = celula(ws, r, c, v, fundo=fundo, strike=sumiu and c <= 11,
                           cor="808080" if sumiu else None, h="center" if c in (1, 3, 5) else "left")
                if c == 1 and primeira:
                    x.font = Font(bold=True, size=10, strike=sumiu, color="808080" if sumiu else None)
                if c == 2 and primeira:
                    x.font = Font(bold=True, size=10, strike=sumiu, color="808080" if sumiu else None)
            if p.get("alerta"):
                ws.cell(row=r, column=10).font = Font(bold=True, color="C00000", size=10)
            if pend.get(p["autos"]):
                ws.cell(row=r, column=11).font = Font(bold=True, color="1F4E79", size=10)
            if primeira:
                x = ws.cell(row=r, column=12)
                x.font = Font(bold=True, color=fg_sit, size=10)
                if bg_sit:
                    x.fill = PatternFill("solid", fgColor=bg_sit)
            ws.cell(row=r, column=COL_OBS).fill = PatternFill("solid", fgColor=COR_EDITAVEL)
            ws.cell(row=r, column=COL_ID).font = Font(size=7, color="FFFFFF")
            texto = max(len(a.get("motivo", "")) if primeira else 0, len(p.get("resumo", "")),
                        len(p.get("atendimentos", "")))
            ws.row_dimensions[r].height = min(max(45, 15 * (texto // 48 + 1)), 400)
            r += 1
    if r == 5:
        celula(ws, 5, 1, "Nenhum agendamento da 9ª nesta janela.", cor="808080", borda=False, wrap=False)
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(CAB) - 1)}{max(r - 1, 4)}"
    return n_ativos


def aba_telefones(wb, agendamentos, periodo):
    ws = wb.create_sheet("Telefone e nº DOL")
    cab = [("Assistido / repr. legal", 32), ("Nº DOL", 11), ("Telefones", 42), ("Modalidade", 13),
           ("Próximo agendamento", 17), ("Processos na 9ª", 30)]
    faixa(ws, "Telefone e nº DOL dos agendados da 9ª",
          f"Janela de {periodo['inicio']} a {periodo['fim']}   |   Dados de contato: uso restrito ao atendimento", len(cab))
    for i, (h, w) in enumerate(cab, 1):
        celula(ws, 4, i, h, bold=True, cor="FFFFFF", fundo=VERDE, h="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    pessoas = {}
    for a in agendamentos:
        if a["situacao"] == "Não aparece mais":
            continue
        p = pessoas.setdefault(a["dol"], {"nome": a["nome"], "obs": a.get("obs_nome", ""), "fones": [], "modos": set(),
                                          "datas": [], "procs": []})
        p["fones"] += [f for f in a.get("fones", []) if f not in p["fones"]]
        p["modos"].add("Presencial" if a["modo"] == "presencial" else "Virtual")
        p["datas"].append(parse_dt(a["data"]))
        p["procs"] += [x["autos"] for x in a["processos"] if x["autos"] not in p["procs"]]
    for r, (dol, p) in enumerate(sorted(pessoas.items(), key=lambda kv: kv[1]["nome"].lower()), 5):
        fundo = GRUPOS[r % 2]
        vals = [p["nome"] + (f" ({p['obs']})" if p["obs"] else ""), dol, "\n".join(p["fones"]) or "sem telefone no cadastro",
                " e ".join(sorted(p["modos"])), min(p["datas"]).strftime("%d/%m/%Y %H:%M"), "\n".join(p["procs"])]
        for c, v in enumerate(vals, 1):
            celula(ws, r, c, v, bold=c == 1, fundo=fundo, h="center" if c in (2, 4, 5) else "left")
        ws.row_dimensions[r].height = max(20, 15 * max(len(p["fones"]), len(p["procs"])))
    ws.freeze_panes = "B5"


def aba_historico(wb, dados, contagem):
    caminho = DADOS / "historico.json"
    hist = json.loads(caminho.read_text(encoding="utf-8")) if caminho.exists() else []
    hist.append({"execucao": datetime.now().strftime("%d/%m/%Y %H:%M"),
                 "janela": f"{dados['periodo']['inicio']} a {dados['periodo']['fim']}",
                 **dados.get("conferidos", {}), **contagem, "mudancas": dados.get("mudancas", "")})
    caminho.write_text(json.dumps(hist, ensure_ascii=False, indent=2), encoding="utf-8")
    ws = wb.create_sheet("Histórico")
    cab = [("Execução", 16), ("Janela", 24), ("Conferidos presencial", 12), ("Conferidos virtual", 12),
           ("Da 9ª presencial", 11), ("Da 9ª virtual", 11), ("Novos", 8), ("Alterados", 9), ("Sumiram", 9),
           ("Mudanças", 80)]
    faixa(ws, "Histórico das conferências", "Mais recente primeiro", len(cab))
    for i, (h, w) in enumerate(cab, 1):
        celula(ws, 4, i, h, bold=True, cor="FFFFFF", fundo=VERDE, h="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    chaves = ["execucao", "janela", "presencial", "virtual", "nove_presencial", "nove_virtual", "novos", "alterados",
              "sumiram", "mudancas"]
    for r, h in enumerate(reversed(hist), 5):
        for c, k in enumerate(chaves, 1):
            celula(ws, r, c, h.get(k, ""), h="left" if c in (2, 10) else "center")
    ws.freeze_panes = "A5"


def aba_como_ler(wb):
    ws = wb.create_sheet("Como ler")
    faixa(ws, "Como ler esta planilha", "Somente leitura do DOL: nada é alterado nos registros originais", 1)
    ws.column_dimensions["A"].width = 125
    linhas = [
        "Fonte: DOL, Visualização geral de agendamentos por unidade (Unidade Ribeirão Preto), tipos Atendimento Presencial - "
        "Cível/Família e Atendimento Virtual - Cível/Família, somente agendamentos ativos.",
        "Janela: sempre de hoje até 7 dias à frente. Todos os dias a janela inteira é conferida de novo, para captar "
        "agendamentos novos, remarcados ou cancelados.",
        "Entram apenas pessoas com processo cuja Localização no Histórico DPESP é '9ª Defensoria da UNIDADE RIBEIRÃO PRETO'. "
        "Cada processo da 9ª ocupa uma linha; linhas da mesma pessoa ficam juntas, com a mesma cor de fundo.",
        "Motivo do agendamento: texto do campo 'Relato do atendido', copiado sem alteração. Quando o sistema não mostra, "
        "aparece a mensagem original do DOL.",
        "Resumo do processo e Últimos atendimentos: resumo feito a partir do histórico do caso e das movimentações no DOL. "
        "Confira sempre no sistema antes de orientar a pessoa.",
        "Pendências no DOL: providências em aberto no controle de prazos para o mesmo processo.",
        "Situação na conferência: Primeira conferência, Novo desde a última conferência, Horário alterado, Sem mudança ou "
        "Não aparece mais (texto riscado: cancelado, remarcado para fora da janela ou deixou de estar ativo).",
        "Observações da equipe (coluna amarela): espaço livre para os estagiários. A atualização diária preserva o que "
        "estiver escrito, lendo a planilha desta pasta antes de gerar a nova versão.",
        "Para o Teams: depois da atualização das 12h30, envie esta planilha para a equipe. Se alguém anotar na versão do "
        "Teams, baixe e salve por cima deste arquivo antes da próxima atualização para não perder as observações.",
        "A aba Telefone e nº DOL reúne o contato dos agendados. Uso restrito ao atendimento.",
    ]
    for i, t in enumerate(linhas, 4):
        celula(ws, i, 1, f"{i - 3}. {t}", borda=False, size=11)
        ws.row_dimensions[i].height = 32


def main():
    saida = Path(sys.argv[1]) if len(sys.argv) > 1 else SAIDA_PADRAO
    saida.parent.mkdir(parents=True, exist_ok=True)
    hoje = date.today()
    dados = json.loads((DADOS / "agendamentos.json").read_text(encoding="utf-8"))
    atuais = [a for a in dados["agendamentos"] if parse_dt(a["data"]).date() >= hoje]
    anterior = conferencia_anterior(hoje)
    sumidos = classificar(atuais, anterior, hoje)
    todos = atuais + sumidos
    notas = carregar_observacoes(saida.parent)
    (DADOS / "observacoes.json").write_text(json.dumps(notas, ensure_ascii=False, indent=2), encoding="utf-8")
    pend = pendencias_dol()

    wb = Workbook()
    wb.remove(wb.active)
    n_pres = aba_agendamentos(wb, "presencial", todos, notas, pend, dados["periodo"])
    n_virt = aba_agendamentos(wb, "virtual", todos, notas, pend, dados["periodo"])
    aba_telefones(wb, todos, dados["periodo"])
    contagem = {"nove_presencial": n_pres, "nove_virtual": n_virt,
                "novos": sum(a["situacao"] == "Novo desde a última conferência" for a in atuais),
                "alterados": sum(a["situacao"] == "Horário alterado" for a in atuais), "sumiram": len(sumidos)}
    aba_historico(wb, dados, contagem)
    aba_como_ler(wb)

    CONFERENCIAS.mkdir(parents=True, exist_ok=True)
    shutil.copy(DADOS / "agendamentos.json", CONFERENCIAS / f"{hoje.isoformat()}.json")
    try:
        wb.save(saida)
        print(f"OK: {saida}")
    except PermissionError:
        alt = saida.with_name(f"{saida.stem} ({datetime.now().strftime('%Y-%m-%d %Hh%M')}){saida.suffix}")
        wb.save(alt)
        print(f"Planilha principal aberta no Excel; salvo em: {alt}")
    print(f"Presencial: {n_pres} | Virtual: {n_virt} | {contagem}")


if __name__ == "__main__":
    main()

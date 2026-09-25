"""Gera a auditoria definitiva dos PAs da 9ª DPE Ribeirão Preto a partir da leitura do DOL (somente leitura).

Uso: python gerar_auditoria.py [caminho_saida.xlsx]
"""
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import base_dol2 as b
from gerar_relacao_eproc import BORDA, VERDE, cabecalho, celula, colorir, faixa, rotulos_barras

SAIDA_PADRAO = Path.home() / "Downloads" / "3. Processos eproc 9ª DPE" / "Auditoria 9ª DPE.xlsx"
CINZA = "F2F2F2"
VERDE_CLARO = "E2EFDA"
VERMELHO = "C00000"
ORDEM_CATEGORIAS = [b.ALIM, b.CURATELA, b.DIV, b.GUARDA, b.FILIACAO, b.CS_ALIM, b.CS_DEMAIS, b.EXEC_ALIM, b.EXEC_TIT,
                    b.AGRAVO, b.OUTROS_TJ, b.INVENT, b.USUCAPIAO, b.MONITORIA, b.FAM_OUTROS, b.CIVEL, b.FISCAL,
                    b.PRISAO, b.SEM_CLASSE]
CUMPRIMENTOS = (b.CS_ALIM, b.CS_DEMAIS, b.EXEC_ALIM)
RESUMOS = [b.PROVIDO, b.PARCIAL, b.NAO_PROVIDO, b.SEGUIMENTO, b.NAO_CONHECIDO, b.PREJUDICADO, b.MISTO,
           "Rejeitados", "Acolhidos", "Não conhecidos", b.AGUARDANDO, b.SEM_RESULTADO]
ASSISTIDO = [b.FAVORAVEL, b.PARC_FAV, b.DESFAVORAVEL, b.PARC_DESF, b.SEM_MERITO,
             "Sem alteração (os recursos das duas partes foram negados)", b.INDETERMINADO, b.AGUARDANDO,
             b.SEM_RESULTADO]
COR_ASSISTIDO = {b.FAVORAVEL: "375623", b.PARC_FAV: "375623", b.DESFAVORAVEL: VERMELHO, b.PARC_DESF: VERMELHO}


def pa_txt(p):
    return p["pa"] or p["id"]


def altura(ws, r, textos_larguras, minimo=18, maximo=150):
    linhas = 1
    for texto, largura in textos_larguras:
        partes = str(texto or "").split("\n")
        linhas = max(linhas, sum(max(1, -(-len(x) // max(1, int(largura * 1.1)))) for x in partes))
    ws.row_dimensions[r].height = min(max(minimo, 13.5 * linhas + 4), maximo)


def titulo_secao(ws, r, texto, c=1):
    celula(ws, r, c, texto, bold=True, cor=VERDE, size=12, borda=False, wrap=False)


def nota(ws, r, texto, c=1):
    celula(ws, r, c, texto, cor="595959", size=9, borda=False, wrap=False)


def cab_tabela(ws, r, titulos, c0=1):
    for i, h in enumerate(titulos):
        celula(ws, r, c0 + i, h, bold=True, cor="FFFFFF", fundo=VERDE, h="center", v_al="center")
    ws.row_dimensions[r].height = 32


def linha_total(ws, r, valores, c0=1, pct=None):
    for i, v in enumerate(valores):
        x = celula(ws, r, c0 + i, v, bold=True, fundo=CINZA, h="left" if i == 0 else "center")
        if isinstance(v, (int, float)) and i:
            x.number_format = "0.0%" if pct == i else "#,##0"


def grafico_barras(ws, titulo, col_val, col_cat, r_ini, r_fim, ancora, cores, pct=True, altura_cm=None):
    g = BarChart()
    g.type = "bar"
    g.style = 10
    g.title = titulo
    g.title.overlay = False
    g.add_data(Reference(ws, min_col=col_val, min_row=r_ini, max_row=r_fim), titles_from_data=True)
    g.set_categories(Reference(ws, min_col=col_cat, min_row=r_ini + 1, max_row=r_fim))
    g.legend = None
    g.x_axis.scaling.orientation = "maxMin"
    g.x_axis.delete = False
    g.y_axis.delete = True
    g.y_axis.majorGridlines = None
    g.gapWidth = 40
    colorir(g.series[0], cores)
    rotulos_barras(g.series[0])
    if pct:
        g.series[0].dLbls.numFmt = "0.0%"
    g.height, g.width = altura_cm or max(7, 0.75 * (r_fim - r_ini) + 3), 21
    ws.add_chart(g, ancora)


def texto_prov(det):
    def campo(rot, prox):
        m = re.search(re.escape(rot) + r":\s*(.*?)\s*(?=" + prox + r"|$)", det)
        return m.group(1).strip() if m else ""
    return {"prazo": campo("Prazo da Providência", "Responsável:")[:10],
            "resp": campo("Responsável", "Providência cumprida:"),
            "cumprida": campo("Providência cumprida", "Urgente"),
            "urgente": campo("Urgente/Alta prioridade", r"[A-ZÁÉÍÓÚÂÊÔÃÕÇa-z0-9]")[:3]}


def foro_vara(p):
    if p["classe_atual"]:
        f = p["classe_atual"]["foro"]
        return f"TJSP, 2º grau (vara de origem: {f})" if b.sg(p["cnj"]) else f
    est = [e for e in p["est"] if e != "Primeira Instância"]
    if est[:1] == ["Segunda Instância"]:
        return "Segunda Instância"
    return " / ".join(est[1:] if len(est) > 1 else est) or "Não informado no DOL"


def sistema(p):
    s = {"E": "eproc", "S": "e-SAJ"}.get(p["sis"], "sem indicação no DOL")
    m = p["mov"]
    if m and not m.get("sem") and m["mg"]:
        d, num = m["mg"][-1]
        extra = "" if num == p["cnj"] else f", número {num}"
        return f"{s}; certidão de remessa ao eproc em {d}{extra}"
    return s


def representados(p):
    if p["nomes"]:
        return "; ".join(dict.fromkeys(p["nomes"]))
    if p["nomes_ficha"]:
        return "; ".join(dict.fromkeys(p["nomes_ficha"]))
    if p["partes"]:
        return "Nenhuma parte marcada como representada. Partes: " + "; ".join(
            f"{n} ({rel})" for n, rel, _ in p["partes"])
    return "Nomes não lidos (correlato sem ficha)" if p["origem"] == "correlato" else "Sem nome na lista do DOL"


def ultimo_andamento(p):
    m = p["mov"]
    if not p["cnj"]:
        return "Sem número de processo"
    if m is None or m.get("sem"):
        return "Nenhum andamento no DOL"
    d, t = m["u3"][0] if m["u3"] else ("", "")
    return f"{d}: {t or '(andamento sem título no DOL)'}"


def texto_situacao(m):
    sit, d = b.situacao_tj(m)
    if sit == b.SEM_MOV:
        return "Sem andamentos no DOL (eproc, outro tribunal ou sem número)"
    return f"{sit} ({'último andamento em ' if sit == b.ANDAMENTO else ''}{d})"


def marcacoes(p):
    m = []
    if p["novo"]:
        m.append(f"Novo na 9ª ({p['novo']})")
    if p["redist"]:
        m.append(f"Redistribuído à 9ª em {p['redist'][0]} (origem: {p['redist'][1]})")
    if p["urgente"]:
        m.append("Urgente")
    if p["sem_prov"]:
        m.append("Sem providência judicial adotada")
    if p["ficha_arq"]:
        m.append("Ficha registra arquivamento")
    if p["acao"] == "Curadoria Especial":
        m.append("Curadoria especial")
    return m


def rotulo_recurso(r):
    num = r["numero"] or "sem número"
    nome = {"Agravo de Instrumento": "Agravo", "Apelação": "Apelação", "Embargos de declaração": "Embargos"}.get(
        r["tipo"], r["tipo"])
    base = f"{nome} {num}" if r["tipo"] != "Apelação" and r["tipo"] != "Embargos de declaração" else nome
    data = f" ({r['data_julg']})" if r["data_julg"] else ""
    return f"{base}: {r['resumo']}{data}; para o assistido: {r['assistido']}"


def vinculos(meta, pas, rec):
    por_pa = {p["pa"]: p for p in pas if p["pa"]}
    por_cnj = defaultdict(list)
    for p in pas:
        if p["cnj"]:
            por_cnj[p["cnj"]].append(p)
    prov, agenda, intim, recs = defaultdict(list), defaultdict(list), defaultdict(list), defaultdict(list)
    for x in meta["prov"]:
        if x[8] in por_pa:
            prov[por_pa[x[8]]["id"]].append(x)
    for x in meta["ag"]:
        if x[8] in por_pa:
            agenda[por_pa[x[8]]["id"]].append(x)
    for x in meta["pend"]:
        for p in por_cnj.get(x[3], []):
            intim[p["id"]].append(x)
    for r in rec:
        if r["principal"]:
            recs[r["principal"]["id"]].append(r)
    return por_pa, por_cnj, prov, agenda, intim, recs


def contagem_categorias(pas):
    cont = {c: {"lista": 0, "corr": 0} for c in ORDEM_CATEGORIAS}
    for p in pas:
        cont[p["categoria"]]["corr" if p["pai"] else "lista"] += 1
    return [(c, v) for c, v in sorted(cont.items(), key=lambda kv: -(kv[1]["lista"] + kv[1]["corr"]))
            if v["lista"] + v["corr"]]


# ------------------------------------------------------------------ abas

def aba_resumo(wb, meta, pas, rec, prov, agenda, intim, arq_meta, arq_lista):
    ws = wb.active
    ws.title = "Resumo"
    faixa(ws, "Auditoria definitiva da 9ª Defensoria  |  Unidade Ribeirão Preto",
          f"Leitura do DOL de {meta['extraido_em']}   |   Somente leitura: nada foi alterado no DOL   |   "
          "Tipo de processo pela classe processual do TJ", 8)
    for col, w in zip("ABCDEFGH", [48, 13, 13, 13, 11, 40, 3, 12]):
        ws.column_dimensions[col].width = w
    total = len(pas)
    lista = sum(1 for p in pas if not p["pai"])
    cnjs = {p["cnj"] for p in pas if p["cnj"]}
    recursos_reais = [r for r in rec if r["recorrente"] != "Não se aplica"]
    titulo_secao(ws, 4, "Visão geral")
    cab_tabela(ws, 5, ["Indicador", "Quantidade"])
    ws.merge_cells("C5:F5")
    celula(ws, 5, 3, "Observação", bold=True, cor="FFFFFF", fundo=VERDE, h="center", v_al="center")
    ind = [
        ("PAs ativos da 9ª", total, "Base de todos os percentuais de tipo"),
        ("    na lista de processos ativos do DOL", lista,
         f"Inclui {sum(1 for p in pas if p['novo'])} novos e {sum(1 for p in pas if p['redist'])} redistribuídos"),
        ("    só como correlatos dentro de outro PA", total - lista,
         "Agravos, cumprimentos de sentença e outros incidentes cadastrados na coluna de correlatos"),
        ("Números de processo distintos", len(cnjs), f"{sum(1 for p in pas if not p['cnj'])} PAs sem número"),
        ("Recursos levantados", len(recursos_reais), "Detalhe nas abas 'Recursos' e 'Recursos (resumo)'"),
        ("    com resultado favorável ou parcialmente favorável ao assistido",
         sum(1 for r in rec if r["assistido"] in b.DEU_CERTO), "Processos principais destacados em verde na aba 'Processos'"),
        ("Providências pendentes", len(meta["prov"]), "Aba 'Providências pendentes'"),
        ("Prazos e audiências na agenda", len(meta["ag"]),
         f"{sum(1 for x in meta['ag'] if x[1].lower().startswith('audiência'))} audiências"),
        ("Intimações recebidas pendentes", len(meta["pend"]), "Aba 'Intimações pendentes'"),
        ("PAs ativos cujo processo já aparece encerrado no TJ",
         sum(1 for p in pas if b.situacao_tj(p["mov"])[0] == b.ENCERRADO),
         "Último andamento é arquivamento ou baixa; aba 'Finalizados recentemente'"),
        ("PAs arquivados da 9ª", (arq_meta or {}).get("total", meta["fontes"]["arquivados"]),
         "Não entram nos percentuais; aba 'Arquivados'"),
        ("    arquivados em 2026", len(arq_lista), "Detalhe na aba 'Finalizados recentemente'"),
    ]
    r = 5
    for rot, val, obs in ind:
        r += 1
        destaque = not rot.startswith("    ")
        celula(ws, r, 1, rot, bold=destaque)
        celula(ws, r, 2, val, bold=destaque, h="center").number_format = "#,##0"
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
        celula(ws, r, 3, obs, cor="595959", size=9, v_al="center")
        for c in range(4, 7):
            ws.cell(row=r, column=c).border = BORDA

    r += 3
    titulo_secao(ws, r, f"Os {total - lista} PAs correlatos, por tipo")
    nota(ws, r + 1, "Correlatos são os PAs cadastrados dentro de outro PA (agravos, cumprimentos de sentença e outros "
                    "incidentes). Cada um tem número próprio no TJ.")
    r += 2
    cab_tabela(ws, r, ["Tipo do correlato", "PAs", "% dos correlatos"])
    corr = Counter(p["categoria"] for p in pas if p["pai"])
    for tipo, q in corr.most_common():
        r += 1
        celula(ws, r, 1, tipo, bold=tipo in (b.AGRAVO, b.CS_ALIM, b.CS_DEMAIS))
        celula(ws, r, 2, q, h="center", bold=True)
        celula(ws, r, 3, q / (total - lista), h="center").number_format = "0.0%"
    r += 1
    linha_total(ws, r, ["Total", total - lista, 1.0], pct=2)
    ws.cell(row=r, column=3).number_format = "0.0%"

    r += 3
    tipos = contagem_categorias(pas)
    titulo_secao(ws, r, f"Tipos de processo dos {total:,} PAs ativos".replace(",", "."))
    nota(ws, r + 1, "Tipo pela classe do processo no TJ (intimação mais recente). Sem intimação, vale a ação do DOL. "
                    "Regras e detalhe na aba 'Tipos (detalhe)'.")
    r += 2
    ini = r
    cab_tabela(ws, r, ["Tipo de processo", "Lista de ativos", "Correlatos", "Total de PAs", "% do total",
                       "Rótulo usado no gráfico"])
    for i, (tipo, c) in enumerate(tipos):
        r += 1
        n = c["lista"] + c["corr"]
        fundo = "F7F7F7" if i % 2 else None
        destaque = tipo in (b.CS_ALIM, b.CS_DEMAIS, b.AGRAVO)
        celula(ws, r, 1, tipo, bold=True, fundo=VERDE_CLARO if destaque else fundo)
        celula(ws, r, 2, c["lista"], h="center", fundo=fundo)
        celula(ws, r, 3, c["corr"], h="center", fundo=fundo)
        celula(ws, r, 4, n, h="center", bold=True, fundo=fundo)
        celula(ws, r, 5, n / total, h="center", bold=True, fundo=fundo).number_format = "0.0%"
        celula(ws, r, 6, f"{tipo}: {n} PA{'s' if n > 1 else ''}", cor="808080", size=9, fundo=fundo, wrap=False)
    fim = r
    r += 1
    linha_total(ws, r, ["Total", lista, total - lista, total, 1.0, ""], pct=4)
    ws.cell(row=r, column=5).number_format = "0.0%"
    cores = [("1F6F43" if t in (b.CS_ALIM, b.CS_DEMAIS, b.AGRAVO) else VERDE) for t, _ in tipos]
    grafico_barras(ws, f"Tipos de processo: % dos {total:,} PAs ativos da 9ª".replace(",", "."), 5, 6, ini, fim,
                   f"H{ini - 2}", cores, altura_cm=15)
    nota(ws, r + 1, "Em verde escuro no gráfico e destacados na tabela: cumprimentos de sentença e agravos de instrumento.")

    r += 4
    titulo_secao(ws, r, "Cumprimentos de sentença e agravos: contagem corrigida")
    nota(ws, r + 1, "A planilha anterior agrupava pela ação cadastrada no DOL. Muitos correlatos herdam a ação do processo "
                    "principal (por exemplo, divórcio), por isso esses incidentes ficavam escondidos.")
    r += 2
    cab_tabela(ws, r, ["Tipo", "Pela ação do DOL (critério anterior)", "Pela classe do TJ (esta auditoria)",
                       "Diferença"])
    ws.row_dimensions[r].height = 44
    antigos = {
        "Cumprimento de sentença (todos)": sum(1 for p in pas if p["acao"] == "Cumprimento de sentença"),
        "Agravo de instrumento": sum(1 for p in pas if p["acao"] == "Agravo de Instrumento"),
    }
    novos = {"Cumprimento de sentença (todos)": sum(1 for p in pas if p["categoria"] in CUMPRIMENTOS),
             "Agravo de instrumento": sum(1 for p in pas if p["categoria"] == b.AGRAVO)}
    for k in antigos:
        r += 1
        celula(ws, r, 1, k, bold=True)
        celula(ws, r, 2, antigos[k], h="center")
        celula(ws, r, 3, novos[k], h="center", bold=True)
        celula(ws, r, 4, novos[k] - antigos[k], h="center", cor="375623", bold=True)
    nota(ws, r + 1, "Pela ação do DOL, só contam como cumprimento os PAs com a ação 'Cumprimento de sentença'; a ação "
                    "'Alimentos - Cumprimento de sentença...' entrava em Alimentos. Os cumprimentos incluem os de "
                    "alimentos, os demais e a execução de alimentos por título extrajudicial.")

    r += 4
    titulo_secao(ws, r, "Cumprimentos de sentença por atuação da Defensoria")
    r += 1
    cab_tabela(ws, r, ["Tipo", "Representação da parte", "Curadoria especial", "Total"])
    for cat in CUMPRIMENTOS:
        r += 1
        cur = sum(1 for p in pas if p["categoria"] == cat and p["acao"] == "Curadoria Especial")
        tot = sum(1 for p in pas if p["categoria"] == cat)
        celula(ws, r, 1, cat, bold=True)
        celula(ws, r, 2, tot - cur, h="center")
        celula(ws, r, 3, cur, h="center")
        celula(ws, r, 4, tot, h="center", bold=True)
    nota(ws, r + 1, "Curadoria especial = PA com a ação 'Curadoria Especial' no DOL (a Defensoria atua como curadora "
                    "especial, em regra de réu citado por edital ou revel).")

    r += 4
    titulo_secao(ws, r, "Situação no TJ dos PAs ativos (andamento mais recente guardado pelo DOL)")
    r += 1
    cab_tabela(ws, r, ["Situação", "PAs", "% do total"])
    cont = Counter(b.situacao_tj(p["mov"])[0] for p in pas)
    for sit in (b.ANDAMENTO, b.ENCERRADO, b.ARQ_PROV, b.TRANSITO, b.MIGRADO, b.SEM_MOV):
        if cont.get(sit):
            r += 1
            celula(ws, r, 1, sit, bold=sit == b.ENCERRADO, cor=VERMELHO if sit == b.ENCERRADO else None)
            celula(ws, r, 2, cont[sit], h="center", bold=True)
            celula(ws, r, 3, cont[sit] / total, h="center").number_format = "0.0%"
    nota(ws, r + 1, "Encerrado no TJ: o último andamento é arquivamento definitivo, envio ao arquivo ou baixa, mas o PA "
                    "continua ativo no DOL. Lista na aba 'Finalizados recentemente'.")
    if arq_meta:
        r += 4
        titulo_secao(ws, r, f"PAs arquivados no DOL em 2026 ({len(arq_lista)} de {arq_meta['total']:,} arquivados)".replace(",", "."))
        r += 1
        cab_tabela(ws, r, ["Mês do arquivamento", "PAs"])
        for mes in sorted(arq_meta["porMes2026"], key=lambda x: x[3:] + x[:2]):
            r += 1
            celula(ws, r, 1, mes, h="center")
            celula(ws, r, 2, arq_meta["porMes2026"][mes], h="center", bold=True)
        nota(ws, r + 1, "Detalhe, com tipo pela classe do TJ e último andamento, na aba 'Finalizados recentemente'.")

    r += 4
    titulo_secao(ws, r, "Marcações que o DOL mostra nos PAs ativos")
    r += 1
    cab_tabela(ws, r, ["Marcação", "PAs"])
    for rot, v in [("Novo na 9ª", sum(1 for p in pas if p["novo"])),
                   ("Redistribuído à 9ª", sum(1 for p in pas if p["redist"])),
                   ("Urgente", sum(1 for p in pas if p["urgente"])),
                   ("Sem providência judicial adotada", sum(1 for p in pas if p["sem_prov"])),
                   ("Curadoria especial (ação no DOL)", sum(1 for p in pas if p["acao"] == "Curadoria Especial")),
                   ("Com providência pendente", len(prov)), ("Com prazo ou audiência na agenda", len(agenda)),
                   ("Com intimação recebida pendente", len(intim))]:
        r += 1
        celula(ws, r, 1, rot)
        celula(ws, r, 2, v, h="center", bold=True)
    ws.freeze_panes = "A4"


def aba_tipos(wb, pas):
    ws = wb.create_sheet("Tipos (detalhe)")
    total = len(pas)
    cols = [("Tipo de processo", 34), ("Classe e assunto no TJ ou ação do DOL", 72), ("Base", 26),
            ("Lista de ativos", 11), ("Correlatos", 11), ("Total de PAs", 11), ("% do total", 10)]
    faixa(ws, "Tipos de processo: de onde vem cada contagem",
          f"{total} PAs   |   Linha verde-clara = total do tipo; abaixo, cada classe do TJ (ou ação do DOL) que entrou nele",
          len(cols))
    cabecalho(ws, 4, cols)
    det = defaultdict(lambda: defaultdict(lambda: {"lista": 0, "corr": 0}))
    for p in pas:
        if p["classe_atual"]:
            chave = (p["classe_atual"]["texto"] + (" (fase de apelação; tipo pela classe anterior)" if p["fase"] else ""),
                     "Classe no TJ")
        elif p["base"].startswith("Sem intimação"):
            chave = ("Número do TJ sem intimação com classe", "Evidência no DOL")
        else:
            chave = (p["acao"] or "(ação não informada)", "Ação do DOL")
        det[p["categoria"]][chave]["corr" if p["pai"] else "lista"] += 1
    r = 4
    for tipo, c in contagem_categorias(pas):
        r += 1
        n = c["lista"] + c["corr"]
        vals = [tipo, f"{len(det[tipo])} classe(s) ou ação(ões)", "", c["lista"], c["corr"], n, n / total]
        for col, v in enumerate(vals, 1):
            x = celula(ws, r, col, v, bold=True, fundo=VERDE_CLARO, h="left" if col < 4 else "center")
        x.number_format = "0.0%"
        for (texto, base), a in sorted(det[tipo].items(), key=lambda kv: -(kv[1]["lista"] + kv[1]["corr"])):
            r += 1
            m = a["lista"] + a["corr"]
            for col, v in enumerate(["", texto, base, a["lista"], a["corr"], m, m / total], 1):
                x = celula(ws, r, col, v, h="left" if col < 4 else "center",
                           cor="7F6000" if base != "Classe no TJ" and col == 3 else None)
            x.number_format = "0.0%"
    r += 1
    linha_total(ws, r, ["Total", "", "", sum(1 for p in pas if not p["pai"]), sum(1 for p in pas if p["pai"]), total, 1.0],
                pct=6)
    ws.cell(row=r, column=7).number_format = "0.0%"
    r += 2
    regras = [
        "Regras de agrupamento (aplicadas à classe do TJ):",
        "Cumprimento de sentença de alimentos: 'Cumprimento de Sentença de Obrigação de Prestar Alimentos', 'Execução de "
        "Alimentos' e qualquer cumprimento (definitivo ou provisório) com assunto 'Alimentos'.",
        "Cumprimento de sentença (demais): as outras classes de cumprimento de sentença, cumprimento provisório e "
        "liquidação de sentença.",
        "Alimentos, Divórcio, Guarda, Filiação e Curatela: classes próprias (por exemplo 'Alimentos - Lei Especial Nº "
        "5.478/68', 'Divórcio Litigioso', 'Guarda de Família', 'Interdição/Curatela') e o 'Procedimento Comum Cível' com o "
        "assunto correspondente (Fixação, Revisão, Dissolução, Guarda, Investigação de Paternidade etc.).",
        "Procedimento comum com assunto de outra matéria: 'Família (outras matérias)' quando a vara é de família; "
        "'Cível (demais ações)' nos demais casos.",
        "Apelação Cível: o número continua o mesmo na apelação; o tipo segue a classe anterior do processo e a aba "
        "'Processos' mostra a fase de apelação.",
        "Sem intimação com classe: número do TJ (2º grau) com peticionamento, ação ou andamento de agravo conta como "
        "agravo; os demais processos usam a ação cadastrada no DOL (coluna 'Base' em amarelo escuro).",
    ]
    for t in regras:
        celula(ws, r, 1, t, borda=False, size=9, cor="404040", bold=t.endswith(":"), wrap=False)
        r += 1
    ws.freeze_panes = "A5"


def aba_processos(wb, pas, por_id, prov, agenda, intim, recs, rec_por_numero):
    ws = wb.create_sheet("Processos")
    cols = [("Nº do PA", 14), ("Número do processo", 26), ("Onde aparece no DOL", 30), ("Tipo de processo", 24),
            ("Classe e assunto no TJ", 34), ("Como foi classificado", 36), ("Ação cadastrada no DOL", 32),
            ("Foro / vara", 32), ("Sistema", 20), ("Representados pela Defensoria", 30),
            ("Situação no TJ (andamento mais recente)", 24), ("Recurso com êxito para o assistido", 13), ("Recursos ligados ao processo (resultado)", 50),
            ("PAs correlatos", 30), ("Andamentos no DOL", 10), ("Último andamento", 34), ("Providências pendentes", 30),
            ("Prazos e audiências", 26), ("Intimações pendentes", 30), ("Marcações", 24)]
    lista = sorted(pas, key=lambda p: (ORDEM_CATEGORIAS.index(p["categoria"]), p["cnj"] or "~", p["id"]))
    faixa(ws, "Todos os PAs ativos da 9ª",
          f"{len(pas)} PAs   |   Linhas em verde: processo com recurso de resultado favorável ou parcialmente "
          "favorável ao assistido   |   Use os filtros do cabeçalho", len(cols))
    cabecalho(ws, 4, cols)
    for n, p in enumerate(lista, 1):
        r = 4 + n
        proprios = recs.get(p["id"], [])
        deste = rec_por_numero.get(p["cnj"], []) if p["cnj"] and b.sg(p["cnj"]) else []
        todos = proprios + [x for x in deste if x not in proprios]
        exito = any(x["assistido"] in b.DEU_CERTO for x in proprios)
        exito_deste = any(x["assistido"] in b.DEU_CERTO for x in deste)
        correlatos = "\n".join(f"PA {pa_txt(por_id[f])} ({por_id[f]['cnj'] or 'sem número'}; {por_id[f]['categoria']})"
                               for f in p["filhos"] if f in por_id)
        onde = "Lista de ativos" if not p["pai"] else \
            f"Correlato do PA {pa_txt(por_id[p['pai']])} ({por_id[p['pai']]['cnj'] or 'sem número'})"
        m = p["mov"]
        classe = p["classe_atual"]["texto"] if p["classe_atual"] else "Sem intimação com classe no DOL"
        if p["fase"]:
            classe += f"\n{p['fase']}"
        vals = [pa_txt(p), p["cnj"] or "(sem número)", onde, p["categoria"], classe, p["base"], p["acao"] or "-",
                foro_vara(p), sistema(p), representados(p), texto_situacao(p["mov"]),
                "Sim" if exito else ("Sim (este é o recurso)" if exito_deste else "-"),
                "\n".join(rotulo_recurso(x) for x in todos) or "-", correlatos or "-",
                m["n"] if m and not m.get("sem") else 0, ultimo_andamento(p),
                "\n".join(f"{x[1]}, inserida em {x[0]}, prazo {texto_prov(x[2])['prazo'] or 'não informado'}"
                          for x in prov.get(p["id"], [])) or "-",
                "\n".join(f"{x[0]}: {x[1]}" for x in agenda.get(p["id"], [])) or "-",
                "\n".join(f"Disponibilizada em {x[0]}, prazo {x[2]} dia(s) ({x[4]})" for x in intim.get(p["id"], [])) or "-",
                "\n".join(marcacoes(p)) or "-"]
        fundo = VERDE_CLARO if exito else ("F7F7F7" if n % 2 == 0 else None)
        for c, v in enumerate(vals, 1):
            celula(ws, r, c, v, fundo=fundo, bold=c in (2, 12) and v not in ("-",),
                   h="center" if c in (1, 12, 15) else "left",
                   cor="375623" if c == 12 and v != "-" else (VERMELHO if c == 11 and v.startswith("Encerrado") else None))
        altura(ws, r, [(vals[i], cols[i][1]) for i in (2, 4, 5, 9, 12, 13, 16, 18)], maximo=130)
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{4 + len(lista)}"


def aba_recursos_resumo(wb, rec):
    ws = wb.create_sheet("Recursos (resumo)")
    faixa(ws, "Recursos da 9ª: tipos, resultados e êxito",
          "Contagens da aba 'Recursos'   |   'Resultado para o assistido' cruza o resultado com quem recorreu", 9)
    for col, w in zip("ABCDEFGHI", [44, 12, 12, 12, 12, 12, 12, 12, 12]):
        ws.column_dimensions[col].width = w
    reais = [r for r in rec if r["recorrente"] != "Não se aplica"]
    tipos = ["Agravo de Instrumento", "Apelação", "Embargos de declaração",
             "Agravo de instrumento (número não informado no DOL)", "Embargos de Declaração Cível", "Agravo Interno Cível",
             "Recurso no TJ (classe não informada no DOL)"]
    r = 4
    titulo_secao(ws, r, "Recursos por tipo e situação")
    r += 1
    cab_tabela(ws, r, ["Tipo de recurso", "Total", "Julgados", "Aguardando julgamento", "Resultado não consta no DOL",
                       "Defensoria recorreu", "Outra parte recorreu", "As duas partes", "Recorrente não identificado"])
    ws.row_dimensions[r].height = 44
    for t in tipos:
        rs = [x for x in reais if x["tipo"] == t]
        if not rs:
            continue
        r += 1
        jul = sum(1 for x in rs if x["resumo"] not in (b.AGUARDANDO, b.SEM_RESULTADO))
        vals = [t, len(rs), jul, sum(1 for x in rs if x["resumo"] == b.AGUARDANDO),
                sum(1 for x in rs if x["resumo"] == b.SEM_RESULTADO), sum(1 for x in rs if x["recorrente"] == b.DPE),
                sum(1 for x in rs if x["recorrente"] == b.OUTRA), sum(1 for x in rs if x["recorrente"] == b.AMBAS),
                sum(1 for x in rs if x["recorrente"] == b.NAO_ID)]
        for c, v in enumerate(vals, 1):
            celula(ws, r, c, v, bold=c in (1, 2), h="left" if c == 1 else "center")
    r += 1
    linha_total(ws, r, ["Total", len(reais), sum(1 for x in reais if x["resumo"] not in (b.AGUARDANDO, b.SEM_RESULTADO)),
                        sum(1 for x in reais if x["resumo"] == b.AGUARDANDO),
                        sum(1 for x in reais if x["resumo"] == b.SEM_RESULTADO),
                        sum(1 for x in reais if x["recorrente"] == b.DPE),
                        sum(1 for x in reais if x["recorrente"] == b.OUTRA),
                        sum(1 for x in reais if x["recorrente"] == b.AMBAS),
                        sum(1 for x in reais if x["recorrente"] == b.NAO_ID)])
    outros = [x for x in rec if x["recorrente"] == "Não se aplica"]
    nota(ws, r + 1, f"Fora da contagem: {len(outros)} feitos originários do Tribunal que não são recursos "
                    f"({', '.join(sorted(set(x['tipo'] for x in outros)))}); estão no fim da aba 'Recursos'.")

    for quem, titulo in ((b.DPE, "Recursos interpostos pela Defensoria"),
                         (b.OUTRA, "Recursos da outra parte (a Defensoria respondeu)")):
        r += 4
        titulo_secao(ws, r, titulo)
        r += 1
        cab_tabela(ws, r, ["Resultado", "Agravos", "Apelações", "Embargos (1º grau)", "Outros", "Total"])
        rs = [x for x in reais if x["recorrente"] == quem]
        for res in RESUMOS:
            sel = [x for x in rs if x["resumo"] == res]
            if not sel:
                continue
            r += 1
            vals = [res, sum(1 for x in sel if x["tipo"] == "Agravo de Instrumento"),
                    sum(1 for x in sel if x["tipo"] == "Apelação"),
                    sum(1 for x in sel if x["tipo"] == "Embargos de declaração"),
                    sum(1 for x in sel if x["tipo"] not in ("Agravo de Instrumento", "Apelação", "Embargos de declaração")),
                    len(sel)]
            for c, v in enumerate(vals, 1):
                celula(ws, r, c, v, bold=c in (1, 6), h="left" if c == 1 else "center")
        r += 1
        linha_total(ws, r, ["Total", sum(1 for x in rs if x["tipo"] == "Agravo de Instrumento"),
                            sum(1 for x in rs if x["tipo"] == "Apelação"),
                            sum(1 for x in rs if x["tipo"] == "Embargos de declaração"),
                            sum(1 for x in rs if x["tipo"] not in ("Agravo de Instrumento", "Apelação",
                                                                   "Embargos de declaração")), len(rs)])
        julg = [x for x in rs if x["resumo"] in (b.PROVIDO, b.PARCIAL, b.NAO_PROVIDO, b.SEGUIMENTO, b.NAO_CONHECIDO,
                                                   "Acolhidos", "Rejeitados", "Não conhecidos")]
        bons = [x for x in julg if x["assistido"] in b.DEU_CERTO or x["assistido"] == b.FAVORAVEL]
        if julg:
            texto = (f"Com julgamento de mérito ou de admissibilidade: {len(julg)}. Resultado favorável ou parcialmente "
                     f"favorável ao assistido: {len(bons)} ({len(bons) / len(julg):.0%}). Não entram nessa conta os "
                     "prejudicados, os que aguardam julgamento e os sem resultado no DOL.")
            nota(ws, r + 1, texto)
        r += 1

    r += 3
    titulo_secao(ws, r, "Resultado para o assistido (todos os recursos)")
    r += 1
    ini = r
    cab_tabela(ws, r, ["Resultado para o assistido", "Recursos", "Rótulo"])
    cont = Counter(x["assistido"] for x in reais)
    for k in ASSISTIDO:
        if cont.get(k):
            r += 1
            celula(ws, r, 1, k, bold=True, cor=COR_ASSISTIDO.get(k))
            celula(ws, r, 2, cont[k], h="center", bold=True)
            celula(ws, r, 3, f"{k}: {cont[k]}", cor="808080", size=9, wrap=False)
    fim = r
    cores = [{b.FAVORAVEL: "375623", b.PARC_FAV: "70AD47", b.DESFAVORAVEL: VERMELHO, b.PARC_DESF: "F4B183"}.get(k, "A6A6A6")
             for k in ASSISTIDO if cont.get(k)]
    grafico_barras(ws, "Resultado dos recursos para o assistido", 2, 3, ini, fim, f"E{ini - 1}", cores, pct=False)
    r += 2
    for t in [
        "Como se lê: se a Defensoria recorreu, provimento é favorável; se a outra parte recorreu, o não provimento é "
        "favorável ao assistido. Quando não se sabe quem recorreu, o resultado para o assistido fica como 'Não "
        "determinável pelo DOL', mesmo que o julgamento seja conhecido.",
        "'Sem julgamento do mérito': recurso prejudicado (perda de objeto, por exemplo por acordo no processo principal).",
    ]:
        nota(ws, r, t)
        r += 1
    ws.freeze_panes = "A4"


def aba_recursos(wb, rec, por_id):
    ws = wb.create_sheet("Recursos")
    cols = [("Tipo de recurso", 22), ("Número do recurso", 26), ("Processo principal", 26), ("Nº do PA", 14),
            ("Onde aparece no DOL", 30), ("Tipo do processo principal", 22), ("Assunto no TJ", 22), ("Matéria", 22),
            ("Câmara", 24), ("Relator(a)", 22), ("Início", 11), ("Quem recorreu", 14),
            ("Como se identificou quem recorreu", 46), ("Resultado (texto do TJ ou do juízo)", 52),
            ("Data do julgamento", 11), ("Resultado resumido", 16), ("Resultado para o assistido", 18),
            ("Trânsito em julgado", 14), ("Observações", 44)]
    ordem_tipo = {"Agravo de Instrumento": 0, "Apelação": 1, "Embargos de declaração": 2,
                  "Agravo de instrumento (número não informado no DOL)": 3}
    lista = sorted(rec, key=lambda x: (x["recorrente"] == "Não se aplica", ordem_tipo.get(x["tipo"], 4),
                                       -int(b.chave(x["inicio"]) or 0), x["numero"]))
    faixa(ws, "Todos os recursos ligados à 9ª",
          f"{len(rec)} linhas   |   Verde: resultado favorável ou parcialmente favorável ao assistido   |   "
          "Cada linha traz a evidência do DOL usada para saber quem recorreu", len(cols))
    cabecalho(ws, 4, cols)
    for n, x in enumerate(lista, 1):
        r = 4 + n
        pr = x["principal"]
        vals = [x["tipo"], x["numero"] or "(não informado)", x["principal_cnj"] or "(não identificado no DOL)",
                pa_txt(pr) if pr else (x["pa_recurso"] or "-"), x["onde"], pr["categoria"] if pr else "-",
                x["assunto"] or "-", x["materia"] or "-", x["orgao"] or "-", x["relator"] or "-", x["inicio"] or "-",
                x["recorrente"], x["evidencia"] or "-", x["resultado"] or "-", x["data_julg"] or "-", x["resumo"],
                x["assistido"], x["transito"] or "-", x["obs"] or "-"]
        fundo = VERDE_CLARO if x["assistido"] in b.DEU_CERTO else ("F7F7F7" if n % 2 == 0 else None)
        for c, v in enumerate(vals, 1):
            celula(ws, r, c, v, fundo=fundo, bold=c in (2, 17), h="center" if c in (11, 12, 15, 16, 17, 18) else "left",
                   cor=COR_ASSISTIDO.get(v) if c == 17 else None)
        altura(ws, r, [(vals[i], cols[i][1]) for i in (4, 12, 13, 18)], maximo=150)
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{4 + len(lista)}"


def aba_cumprimentos(wb, meta, pas, por_id):
    ws = wb.create_sheet("Cumprimentos de sentença")
    sel = [p for p in pas if p["categoria"] in CUMPRIMENTOS]
    cols = [("Nº do PA", 14), ("Número do processo", 26), ("Tipo", 30), ("Classe e assunto no TJ", 40),
            ("Atuação da Defensoria", 20), ("Onde aparece no DOL", 30), ("Processo principal e tipo", 40),
            ("Foro / vara", 32), ("Representados pela Defensoria", 32), ("Último andamento", 36)]
    faixa(ws, "Cumprimentos de sentença e execuções de alimentos da 9ª",
          f"{len(sel)} PAs   |   Classificação pela classe do TJ   |   Inclui os incidentes cadastrados como correlatos",
          len(cols))
    cabecalho(ws, 4, cols)
    lista = sorted(sel, key=lambda p: (CUMPRIMENTOS.index(p["categoria"]), p["acao"] == "Curadoria Especial", p["cnj"]))
    for n, p in enumerate(lista, 1):
        r = 4 + n
        pai = por_id.get(p["pai"]) if p["pai"] else None
        vals = [pa_txt(p), p["cnj"] or "(sem número)", p["categoria"],
                p["classe_atual"]["texto"] if p["classe_atual"] else f"Sem classe no TJ; {p['base']}",
                "Curadoria especial" if p["acao"] == "Curadoria Especial" else "Representação da parte",
                "Lista de ativos" if not pai else f"Correlato do PA {pa_txt(pai)}",
                f"{pai['cnj'] or 'sem número'} ({pai['categoria']})" if pai else "É o próprio processo da lista",
                foro_vara(p), representados(p), ultimo_andamento(p)]
        fundo = "F7F7F7" if n % 2 == 0 else None
        for c, v in enumerate(vals, 1):
            celula(ws, r, c, v, fundo=fundo, bold=c == 2, h="center" if c == 1 else "left")
        altura(ws, r, [(vals[i], cols[i][1]) for i in (3, 6, 8, 9)], maximo=90)
    r = 6 + len(lista)
    titulo_secao(ws, r, "Peticionamentos da Defensoria ligados a cumprimento de sentença (DOL, status Enviado)")
    r += 1
    cab_tabela(ws, r, ["Ato no peticionamento", "Quantidade"])
    cont = Counter(x[3] for p in pas for x in p["pet"] if x[5] == "Enviado" and "umprimento" in x[3])
    for ato, q in cont.most_common():
        r += 1
        celula(ws, r, 1, ato)
        celula(ws, r, 2, q, h="center", bold=True)
    nota(ws, r + 1, "Os pedidos de início de cumprimento geram número próprio no TJ; os que já têm PA aparecem na lista acima.")
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{4 + len(lista)}"


def aba_finalizados(wb, pas, por_id, arq_meta, arq_lista):
    ws = wb.create_sheet("Finalizados recentemente")
    cols = [("Nº do PA", 14), ("Número do processo", 26), ("Data do arquivamento no DOL", 13), ("Tipo de processo", 26),
            ("Classe e assunto no TJ", 36), ("Ação cadastrada no DOL", 32), ("Foro / vara", 32), ("Sistema", 10),
            ("Situação no TJ (andamento mais recente)", 30), ("Último andamento de sentença, trânsito, arquivamento, baixa ou remessa ao eproc", 36),
            ("Correlatos do PA", 30)]
    faixa(ws, "Processos finalizados recentemente",
          "Parte 1: PAs arquivados no DOL em 2026   |   Parte 2: PAs ainda ativos cujo processo já aparece encerrado no TJ",
          len(cols))
    r = 4
    if arq_lista:
        titulo_secao(ws, r, f"Parte 1. PAs arquivados no DOL em 2026: {len(arq_lista)} (leitura de "
                            f"{arq_meta['extraido_em']}; em verde, os arquivados a partir de 25/06/2026)")
        r += 1
        cabecalho(ws, r, cols)
        ini = r
        lista = sorted(arq_lista, key=lambda p: (b.chave(p["data_arq"]), p["cnj"]), reverse=True)
        for n, p in enumerate(lista, 1):
            r += 1
            m = p["mov"]
            fim = m.get("fim") if m and not m.get("sem") else None
            if fim and (fim[1].startswith("Conclusos") or "Impugnação" in fim[1] or "Desarquivado" in fim[1]):
                fim = None
            vals = [pa_txt(p), p["cnj"] or "(sem número)", p["data_arq"], p["categoria"],
                    p["classe_atual"]["texto"] if p["classe_atual"] else "Sem intimação com classe; " + p["base"],
                    p["acao"] or "-", foro_vara(p), {"E": "eproc", "S": "e-SAJ"}.get(p["sis"], "sem indicação"),
                    texto_situacao(m), f"{fim[0]}: {fim[1]}" if fim else "-",
                    "\n".join(f"{c[0] or 'sem número'} (PA {c[1]})" for c in p["cor"]) or "-"]
            recente = b.chave(p["data_arq"]) >= "20260625"
            fundo = VERDE_CLARO if recente else ("F7F7F7" if n % 2 == 0 else None)
            for c, v in enumerate(vals, 1):
                celula(ws, r, c, v, fundo=fundo, bold=c == 2, h="center" if c in (1, 3, 8) else "left")
            altura(ws, r, [(vals[i], cols[i][1]) for i in (4, 5, 6, 8, 9, 10)], maximo=80)
        ws.auto_filter.ref = f"A{ini}:{get_column_letter(len(cols))}{r}"
        r += 2
        titulo_secao(ws, r, "Tipos dos PAs arquivados em 2026")
        r += 1
        cab_tabela(ws, r, ["Tipo de processo", "PAs", "% dos arquivados em 2026"])
        for tipo, q in Counter(p["categoria"] for p in arq_lista).most_common():
            r += 1
            celula(ws, r, 1, tipo, bold=True)
            celula(ws, r, 2, q, h="center", bold=True)
            celula(ws, r, 3, q / len(arq_lista), h="center").number_format = "0.0%"
        r += 3
    enc = [p for p in pas if b.situacao_tj(p["mov"])[0] in (b.ENCERRADO, b.TRANSITO)]
    titulo_secao(ws, r, f"Parte 2. PAs ativos no DOL cujo processo já aparece encerrado ou com trânsito no TJ: {len(enc)}")
    r += 1
    cols2 = [("Nº do PA", 14), ("Número do processo", 26), ("Onde aparece no DOL", 13), ("Tipo de processo", 26),
             ("Classe e assunto no TJ", 36), ("Ação cadastrada no DOL", 32), ("Foro / vara", 32), ("Sistema", 10),
             ("Situação no TJ (andamento mais recente)", 30), ("Três andamentos mais recentes", 36),
             ("Correlatos do PA", 30)]
    for i, (h, _) in enumerate(cols2, 1):
        celula(ws, r, i, h, bold=True, cor="FFFFFF", fundo=VERDE, h="center", v_al="center")
    ws.row_dimensions[r].height = 44
    lista2 = sorted(enc, key=lambda p: b.chave(b.situacao_tj(p["mov"])[1]), reverse=True)
    for n, p in enumerate(lista2, 1):
        r += 1
        m = p["mov"]
        vals = [pa_txt(p), p["cnj"], "Lista de ativos" if not p["pai"] else f"Correlato do PA {pa_txt(por_id[p['pai']])}",
                p["categoria"], p["classe_atual"]["texto"] if p["classe_atual"] else "Sem intimação com classe",
                p["acao"] or "-", foro_vara(p), {"E": "eproc", "S": "e-SAJ"}.get(p["sis"], "sem indicação"),
                texto_situacao(m), "\n".join(f"{d}: {t or '(sem título)'}" for d, t in m["u3"]),
                "\n".join(f"{por_id[f]['cnj'] or 'sem número'} (PA {pa_txt(por_id[f])})"
                          for f in p["filhos"] if f in por_id) or "-"]
        fundo = "F7F7F7" if n % 2 == 0 else None
        for c, v in enumerate(vals, 1):
            celula(ws, r, c, v, fundo=fundo, bold=c == 2, h="center" if c in (1, 8) else "left")
        altura(ws, r, [(vals[i], cols2[i][1]) for i in (2, 4, 5, 9, 10)], maximo=80)
    ws.freeze_panes = "C4"


def aba_providencias(wb, meta, por_pa):
    ws = wb.create_sheet("Providências pendentes")
    cols = [("Inserida em", 12), ("Tipo de providência", 22), ("Prazo", 12), ("Urgente", 9), ("Cumprida", 9),
            ("Número do processo", 26), ("Nº do PA", 14), ("Tipo do processo", 26), ("Onde aparece no DOL", 30),
            ("Responsável (DOL)", 30), ("Inserida por (DOL)", 30)]
    faixa(ws, "Providências pendentes da 9ª",
          f"{len(meta['prov'])} providências   |   Ordenadas pelo prazo   |   O texto livre da providência não foi copiado",
          len(cols))
    cabecalho(ws, 4, cols)
    linhas = []
    for x in meta["prov"]:
        t = texto_prov(x[2])
        p = por_pa.get(x[8])
        linhas.append((t, x, p))
    linhas.sort(key=lambda z: (b.chave(z[0]["prazo"]) if re.fullmatch(r"\d\d/\d\d/\d{4}", z[0]["prazo"]) else "99999999"))
    for n, (t, x, p) in enumerate(linhas, 1):
        onde = "-" if not p else ("Lista de ativos" if not p["pai"] else f"Correlato do PA {p['pai']}")
        vals = [x[0], x[1], t["prazo"] or "Sem prazo no DOL", t["urgente"] or "-", t["cumprida"] or "-", x[5] or "-",
                x[8], p["categoria"] if p else "PA não localizado entre os ativos", onde, x[10] or "-", x[11] or "-"]
        fundo = "F7F7F7" if n % 2 == 0 else None
        for c, v in enumerate(vals, 1):
            celula(ws, 4 + n, c, v, fundo=fundo, h="center" if c in (1, 3, 4, 5, 7) else "left", bold=c == 3)
        if t["urgente"].startswith("Sim"):
            ws.cell(row=4 + n, column=4).font = Font(bold=True, color=VERMELHO, size=10, name="Calibri")
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{4 + len(linhas)}"


def data_agenda(s):
    m = re.match(r"(\d\d)/(\d\d)/(\d{4})(?: às (\d\d):(\d\d))?", s)
    return (m.group(3) + m.group(2) + m.group(1) + (m.group(4) or "00") + (m.group(5) or "00")) if m else "9"


def aba_agenda(wb, meta, por_pa):
    ws = wb.create_sheet("Prazos e audiências")
    cols = [("Data mostrada pelo DOL", 20), ("Compromisso", 50), ("Número do processo", 26), ("Nº do PA", 14),
            ("Tipo do processo", 26), ("Onde aparece no DOL", 30)]
    faixa(ws, "Prazos e audiências pendentes na agenda da 9ª",
          f"{len(meta['ag'])} compromissos   |   Nos itens 'Prazo - Providências' a data exibida nem sempre é o prazo "
          "final; o prazo de cada providência está na aba 'Providências pendentes'", len(cols))
    cabecalho(ws, 4, cols)
    rows = sorted(meta["ag"], key=lambda x: data_agenda(x[0]))
    for n, x in enumerate(rows, 1):
        p = por_pa.get(x[8])
        aud = x[1].lower().startswith("audiência")
        vals = [x[0], x[1], x[4] or "-", x[8], p["categoria"] if p else "PA não localizado entre os ativos",
                "-" if not p else ("Lista de ativos" if not p["pai"] else f"Correlato do PA {p['pai']}")]
        fundo = VERDE_CLARO if aud else ("F7F7F7" if n % 2 == 0 else None)
        for c, v in enumerate(vals, 1):
            celula(ws, 4 + n, c, v, fundo=fundo, h="center" if c in (1, 4) else "left", bold=aud and c in (1, 2))
    nota(ws, 6 + len(rows), "Linhas em verde: audiências.")
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{4 + len(rows)}"


def aba_intimacoes(wb, meta, por_cnj):
    ws = wb.create_sheet("Intimações pendentes")
    cols = [("Disponibilização", 13), ("Intimação", 12), ("Prazo (dias, como no DOL)", 11), ("Número do processo", 26),
            ("Classe / assunto (DOL)", 40), ("Foro / vara", 36), ("Recebida por (DOL)", 26),
            ("PA(s) da 9ª com este número", 16), ("Tipo do processo", 26)]
    faixa(ws, "Intimações recebidas pendentes da 9ª",
          f"{len(meta['pend'])} intimações   |   Lista de intimações recebidas com situação pendente", len(cols))
    cabecalho(ws, 4, cols)
    rows = sorted(meta["pend"], key=lambda x: (b.chave(x[1]), x[3]))
    for n, x in enumerate(rows, 1):
        ps = por_cnj.get(x[3], [])
        vals = list(x[:7]) + ["\n".join(pa_txt(p) for p in ps) or "Nenhum PA ativo",
                              "\n".join(dict.fromkeys(p["categoria"] for p in ps)) or "-"]
        fundo = "F7F7F7" if n % 2 == 0 else None
        for c, v in enumerate(vals, 1):
            celula(ws, 4 + n, c, v, fundo=fundo, h="center" if c in (1, 2, 3, 8) else "left")
        altura(ws, 4 + n, [(vals[4], 40), (vals[5], 36)])
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{4 + len(rows)}"


def pontos_atencao(meta, pas, por_id, por_cnj, prov, agenda, intim, rec):
    pontos = []

    def add(titulo, p, detalhe):
        pontos.append((titulo, pa_txt(p) if p else "-", (p["cnj"] or "(sem número)") if p else detalhe[0],
                       p["categoria"] if p else "-",
                       ("Lista de ativos" if not p["pai"] else f"Correlato do PA {pa_txt(por_id[p['pai']])}") if p else "-",
                       detalhe if p else detalhe[1]))

    for p in pas:
        if p["classe_atual"] and p["acao"] not in ("", "Curadoria Especial"):
            pela_acao = b.categoria_acao(p["acao"])
            if pela_acao != p["categoria"] and pela_acao not in (b.SEM_CLASSE, b.CIVEL):
                add("Ação cadastrada no DOL indica outro tipo que a classe atual no TJ", p,
                    f"DOL: {p['acao']} (seria {pela_acao}); TJ: {p['classe_atual']['texto']}")
    for cnj, ps in sorted(por_cnj.items()):
        if len(ps) > 1:
            for p in ps:
                add("Mesmo número de processo em dois PAs ativos", p,
                    "Também no PA " + ", ".join(pa_txt(q) for q in ps if q is not p))
    for p in pas:
        if not p["cnj"]:
            add("PA ativo sem número de processo", p, f"Classificado pela ação do DOL: {p['acao'] or 'não informada'}")
        elif ".8.26." not in p["cnj"]:
            add("Processo de outro tribunal (número não é do TJSP)", p, "O DOL não guarda andamentos deste processo")
        elif not p["classe_atual"]:
            add("Número sem intimação com classe no DOL", p, p["base"])
    for x in rec:
        if x["onde"].startswith("Só nas intimações") and x["tipo"] != "Embargos de declaração":
            add_linha = (x["numero"], f"{x['tipo']}: {x['onde']}. Resultado: {x['resumo']}")
            pontos.append(("Recurso nas intimações da 9ª sem PA ativo", "-", x["numero"], x["tipo"], "-",
                           add_linha[1]))
        if x["tipo"].startswith("Agravo de instrumento (número"):
            add("Agravo registrado no 1º grau sem número no DOL", x["principal"], x["obs"]) if x["principal"] else None
        if x["recorrente"] == b.NAO_ID and x["resumo"] not in (b.SEM_RESULTADO,) and x["principal"] and \
                x["tipo"] == "Agravo de Instrumento":
            add("Agravo julgado ou em curso sem registro de quem recorreu", x["principal"],
                f"Agravo {x['numero']}: {x['resumo']}; {x['evidencia']}")
    for p in pas:
        if p["ficha_arq"]:
            add("Correlato cuja ficha registra arquivamento", p, "A ficha indica arquivamento, mas ele segue como correlato")
        m = p["mov"]
        if m and not m.get("sem") and m["mg"] and p["sis"] != "E":
            add("Certidão de remessa ao eproc, mas o DOL não indica eproc", p, sistema(p))
        if p["partes"] and not any(dpe for _, _, dpe in p["partes"]):
            add("Nenhuma parte marcada como 'Representado por Defensoria' na ficha", p, representados(p))
        if p["pai"] and (p["id"] in prov or p["id"] in agenda or p["id"] in intim):
            add("Pendência em PA que só aparece como correlato", p,
                f"{len(prov.get(p['id'], []))} providência(s), {len(agenda.get(p['id'], []))} item(ns) de agenda, "
                f"{len(intim.get(p['id'], []))} intimação(ões)")
        if b.situacao_tj(p["mov"])[0] == b.ENCERRADO:
            add("Processo encerrado no TJ, mas o PA segue ativo no DOL", p, texto_situacao(p["mov"]))
        if p["sem_prov"]:
            add("Marcado no DOL como 'sem providência judicial adotada'", p, "; ".join(marcacoes(p)))
        if p["urgente"]:
            add("Marcado como urgente", p, "; ".join(marcacoes(p)))
        if p["novo"] or p["redist"]:
            add("Entrada recente na 9ª (novo ou redistribuído)", p, "; ".join(marcacoes(p)))
    return pontos


def aba_pontos(wb, pontos):
    ws = wb.create_sheet("Pontos de atenção")
    cols = [("Ponto de atenção", 46), ("Nº do PA", 14), ("Número do processo", 26), ("Tipo", 24),
            ("Onde aparece no DOL", 30), ("Detalhe (dados do DOL)", 80)]
    faixa(ws, "Pontos de atenção encontrados na leitura",
          "Somente fatos que constam no DOL; cada ponto é uma indicação para conferência, não uma conclusão", len(cols))
    ws.column_dimensions["A"].width = 46
    titulo_secao(ws, 4, "Resumo")
    cab_tabela(ws, 5, ["Ponto de atenção", "Linhas"])
    cont = Counter(p[0] for p in pontos)
    r = 5
    for t in dict.fromkeys(p[0] for p in pontos):
        r += 1
        celula(ws, r, 1, t)
        celula(ws, r, 2, cont[t], h="center", bold=True)
    r += 2
    titulo_secao(ws, r, "Detalhe (filtre pela coluna A)")
    r += 1
    cab = r
    cabecalho(ws, cab, cols)
    for n, linha in enumerate(pontos, 1):
        r += 1
        fundo = "F7F7F7" if n % 2 == 0 else None
        for c, v in enumerate(linha, 1):
            celula(ws, r, c, v, fundo=fundo, h="center" if c == 2 else "left", bold=c == 3)
        altura(ws, r, [(linha[0], 46), (linha[4], 30), (linha[5], 80)], maximo=110)
    ws.auto_filter.ref = f"A{cab}:{get_column_letter(len(cols))}{r}"
    ws.freeze_panes = "A4"


def aba_arquivados(wb, meta):
    ws = wb.create_sheet("Arquivados")
    arq = meta["arq"]
    total = meta["fontes"]["arquivados"]
    faixa(ws, "PAs arquivados da 9ª",
          f"{total:,} PAs na lista de arquivados do DOL   |   Apenas totais: os andamentos desses processos não foram "
          "lidos".replace(",", "."), 8)
    for col, w in zip("ABCDEFGH", [44, 12, 12, 3, 12, 12, 12, 12]):
        ws.column_dimensions[col].width = w
    titulo_secao(ws, 4, "Por ano do arquivamento")
    cab_tabela(ws, 5, ["Ano", "PAs", "%"])
    r = 5
    for ano in sorted(arq["porAno"]):
        r += 1
        celula(ws, r, 1, ano, h="center")
        celula(ws, r, 2, arq["porAno"][ano], h="center", bold=True)
        celula(ws, r, 3, arq["porAno"][ano] / total, h="center").number_format = "0.0%"
    fim = r
    r += 1
    linha_total(ws, r, ["Total", sum(arq["porAno"].values()), 1.0], pct=2)
    ws.cell(row=r, column=3).number_format = "0.0%"
    g = BarChart()
    g.type = "col"
    g.style = 10
    g.title = "PAs arquivados por ano do arquivamento"
    g.add_data(Reference(ws, min_col=2, min_row=5, max_row=fim), titles_from_data=True)
    g.set_categories(Reference(ws, min_col=1, min_row=6, max_row=fim))
    g.legend = None
    g.y_axis.delete = True
    g.y_axis.majorGridlines = None
    g.x_axis.delete = False
    colorir(g.series[0], [VERDE] * (fim - 5))
    rotulos_barras(g.series[0])
    g.height, g.width = 8, 18
    ws.add_chart(g, "E4")
    r += 3
    titulo_secao(ws, r, "Por sistema (botão de pasta digital na lista de arquivados)")
    r += 1
    cab_tabela(ws, r, ["Sistema", "PAs", "%"])
    for k, nome in (("S", "e-SAJ"), ("E", "eproc"), ("-", "Sem botão de pasta digital na lista")):
        if arq["porSis"].get(k):
            r += 1
            celula(ws, r, 1, nome)
            celula(ws, r, 2, arq["porSis"][k], h="center", bold=True)
            celula(ws, r, 3, arq["porSis"][k] / total, h="center").number_format = "0.0%"
    r += 3
    titulo_secao(ws, r, "Por tipo (pela ação do DOL, pois a classe do TJ dos arquivados não foi lida)")
    r += 1
    cab_tabela(ws, r, ["Tipo de processo", "PAs", "%"])
    por_tipo = Counter()
    for acao, n in arq["porAcao"].items():
        por_tipo[b.categoria_acao(acao) if acao != "Curadoria Especial" else "Curadoria especial (matéria não indicada)"] += n
    for tipo, n in por_tipo.most_common():
        r += 1
        celula(ws, r, 1, tipo)
        celula(ws, r, 2, n, h="center", bold=True)
        celula(ws, r, 3, n / total, h="center").number_format = "0.0%"
    r += 3
    titulo_secao(ws, r, "Arquivados que o DOL indica como eproc")
    r += 1
    cab_tabela(ws, r, ["Número do processo", "Nº do PA", "Arquivado em"])
    for x in arq["eproc"]:
        r += 1
        celula(ws, r, 1, x[2])
        celula(ws, r, 2, x[1], h="center")
        celula(ws, r, 3, x[3], h="center")
    ws.freeze_panes = "A4"


def aba_contagens(wb, pas):
    ws = wb.create_sheet("Outras contagens")
    total = len(pas)
    faixa(ws, "Outras contagens dos PAs ativos", f"{total} PAs ativos", 7)
    for col, w in zip("ABCDEFG", [70, 10, 10, 3, 30, 10, 10]):
        ws.column_dimensions[col].width = w
    titulo_secao(ws, 4, "Por foro e vara")
    cab_tabela(ws, 5, ["Foro / vara", "PAs", "%"])
    r = 5
    for fv, n in Counter(foro_vara(p) for p in pas).most_common():
        r += 1
        celula(ws, r, 1, fv)
        celula(ws, r, 2, n, h="center", bold=True)
        celula(ws, r, 3, n / total, h="center").number_format = "0.0%"
    titulo_secao(ws, 4, "Por ano do PA", c=5)
    cab_tabela(ws, 5, ["Ano do PA", "PAs", "%"], c0=5)
    r2 = 5
    anos = Counter(p["pa"].split("/")[1] for p in pas if "/" in p["pa"])
    for ano in sorted(anos):
        r2 += 1
        celula(ws, r2, 5, ano, h="center")
        celula(ws, r2, 6, anos[ano], h="center", bold=True)
        celula(ws, r2, 7, anos[ano] / total, h="center").number_format = "0.0%"
    ws.freeze_panes = "A4"


def aba_metodo(wb, meta, pas, rec, n_validos, n_numeros, arq_meta=None):
    ws = wb.create_sheet("Como foi feito")
    faixa(ws, "Como a auditoria foi feita", "Consulta feita somente em modo leitura no DOL", 1)
    ws.column_dimensions["A"].width = 140
    f = meta["fontes"]
    linhas = [
        f"Fonte: DOL, Acompanhamento de Processo da 9ª Defensoria da UNIDADE RIBEIRÃO PRETO, lido em {meta['extraido_em']}: "
        f"lista de PAs ativos ({f['ativo']}), os {f['corr']} PAs correlatos com as fichas, os peticionamentos de cada PA, "
        "as movimentações do e-SAJ de cada número, o histórico de intimações recebidas "
        f"({meta['intTot']['n']:,} intimações de {meta['intTot']['distintos']:,} números), as intimações pendentes, as "
        "providências pendentes, a agenda e os totais dos arquivados.".replace(",", "."),
        "Tipo de processo: pela classe processual do TJ que aparece na intimação mais recente de cada número (coluna "
        "'Classe / Assunto principal' das intimações). É o que corrige a contagem de cumprimentos de sentença e agravos: "
        "no DOL, muitos correlatos repetem a ação do processo principal. Quando o número não tem intimação, o tipo vem da "
        "ação cadastrada no DOL, e isso fica indicado na coluna 'Como foi classificado'. As regras estão na aba 'Tipos "
        "(detalhe)'.",
        "Recursos: (1) todo número do 2º grau ligado à 9ª, seja correlato de um PA, seja número que só aparece nas "
        "intimações; (2) as apelações, identificadas nas movimentações do processo (razões juntadas, remessa ao TJ, "
        "entrada de recursos, retorno dos autos e decisão de 2ª instância) e nas intimações da classe Apelação; (3) os "
        "embargos de declaração juntados no 1º grau; (4) registros de agravo no 1º grau sem agravo com número ligado ao PA.",
        "Resultado: texto do julgamento como está no DOL ('Julgado virtualmente', decisão monocrática ou decisão de 2ª "
        "instância juntada no processo). Quando o DOL só mostra o retorno dos autos do TJ, o resultado fica como 'não "
        "consta no DOL'. Nada foi deduzido fora desses textos.",
        "Quem recorreu: (a) peticionamento da Defensoria com o número do recurso como petição inicial de 2º grau; (b) "
        "contraminuta ou contrarrazões da Defensoria no PA do recurso; (c) protocolo das razões de apelação ou dos "
        "embargos juntados no processo igual ao protocolo de um peticionamento da Defensoria; (d) nome ou iniciais do "
        "recorrente citados no texto do TJ que conferem com pessoa representada pela Defensoria no PA; (e) providência "
        "pendente que manda fazer contraminuta de um recurso ainda não julgado. Sem nenhuma dessas provas, fica 'Não "
        "identificado no DOL'. A coluna 'Como se identificou quem recorreu' mostra a prova usada em cada linha.",
        "Resultado para o assistido: recurso da Defensoria provido = favorável; provido em parte = parcialmente "
        "favorável; negado, não conhecido ou com seguimento negado = desfavorável. Recurso da outra parte: o inverso. "
        "Prejudicado = sem julgamento do mérito. Se não se sabe quem recorreu, 'Não determinável pelo DOL'. Os processos "
        "principais com pelo menos um recurso favorável ou parcialmente favorável estão em verde na aba 'Processos'.",
        f"Conferências: {n_validos} de {n_numeros} números de processo citados nesta planilha têm o dígito verificador "
        "correto pela regra do CNJ. A exportação do DOL foi transferida em 18 blocos, cada um conferido por soma de "
        "verificação, e o arquivo inteiro conferiu com a soma calculada no próprio DOL.",
        "Finalizados recentemente: (1) PAs da lista de arquivados do DOL com data de arquivamento em 2026, lidos em "
        f"{arq_meta['extraido_em'] if arq_meta else 'leitura própria'}, com a classe do TJ pelas intimações e o último "
        "andamento do e-SAJ; (2) PAs ainda ativos cujo andamento mais recente guardado pelo DOL é arquivamento "
        "definitivo, envio ao arquivo ou baixa (ou trânsito em julgado com o processo seguindo).",
        "Nomes: só os nomes que o DOL mostra como representados pela Defensoria. Nenhum CPF, RG, endereço ou contato foi "
        "copiado. O texto livre das providências e da agenda não foi copiado.",
        "Limites: o DOL não guarda andamentos do eproc nem de outros tribunais, e a lista de peticionamentos só mostra o "
        "que foi protocolado pelo DOL. Por isso parte dos recursos antigos fica sem recorrente ou sem resultado. Os "
        "pontos de atenção indicam o que vale conferir; não afirmam que haja erro.",
        "Nenhum dado foi incluído, alterado ou excluído no DOL: foram feitas apenas consultas de leitura.",
    ]
    for i, texto in enumerate(linhas, 4):
        celula(ws, i, 1, f"{i - 3}. {texto}", borda=False, size=11)
        ws.row_dimensions[i].height = 16 * (len(texto) // 125 + 1) + 6


def numeros_citados(meta, pas, rec):
    nums = {p["cnj"] for p in pas if p["cnj"]}
    nums |= {x["numero"] for x in rec if x["numero"]} | {x["principal_cnj"] for x in rec if x["principal_cnj"]}
    nums |= {x[5] for x in meta["prov"] if x[5]} | {x[4] for x in meta["ag"] if x[4]} | {x[3] for x in meta["pend"]}
    return {n for n in nums if b.RX_CNJ.match(n)}


def main():
    saida = Path(sys.argv[1]) if len(sys.argv) > 1 else SAIDA_PADRAO
    meta, pas, por_id, classes, mov = b.carregar()
    arq_meta, arq_lista = b.carregar_arquivados()
    rec = b.recursos(meta, pas, por_id, classes, mov)
    por_pa, por_cnj, prov, agenda, intim, recs = vinculos(meta, pas, rec)
    rec_por_numero = defaultdict(list)
    for x in rec:
        if x["numero"]:
            rec_por_numero[x["numero"]].append(x)
    nums = numeros_citados(meta, pas, rec) | {p["cnj"] for p in arq_lista if b.RX_CNJ.match(p["cnj"] or "")}
    validos = sum(1 for n in nums if b.cnj_valido(n))
    pontos = pontos_atencao(meta, pas, por_id, por_cnj, prov, agenda, intim, rec)

    wb = Workbook()
    aba_resumo(wb, meta, pas, rec, prov, agenda, intim, arq_meta, arq_lista)
    aba_tipos(wb, pas)
    aba_processos(wb, pas, por_id, prov, agenda, intim, recs, rec_por_numero)
    aba_recursos_resumo(wb, rec)
    aba_recursos(wb, rec, por_id)
    aba_cumprimentos(wb, meta, pas, por_id)
    aba_finalizados(wb, pas, por_id, arq_meta, arq_lista)
    aba_providencias(wb, meta, por_pa)
    aba_agenda(wb, meta, por_pa)
    aba_intimacoes(wb, meta, por_cnj)
    aba_pontos(wb, pontos)
    aba_arquivados(wb, meta)
    aba_contagens(wb, pas)
    aba_metodo(wb, meta, pas, rec, validos, len(nums), arq_meta)

    try:
        wb.save(saida)
        print(f"OK: {saida}")
    except PermissionError:
        alt = saida.with_name(f"{saida.stem} ({datetime.now().strftime('%Y-%m-%d %Hh%M')}){saida.suffix}")
        wb.save(alt)
        print(f"Planilha principal aberta no Excel; salvo em: {alt}")
    print(f"PAs: {len(pas)} | recursos: {len(rec)} | números válidos: {validos}/{len(nums)}")
    print("Tipos:", [(t, c["lista"] + c["corr"]) for t, c in contagem_categorias(pas)])
    print("Pontos:", dict(Counter(p[0] for p in pontos)))
    print("Vínculos: prov", sum(map(len, prov.values())), "agenda", sum(map(len, agenda.values())),
          "intim", sum(map(len, intim.values())))


if __name__ == "__main__":
    main()

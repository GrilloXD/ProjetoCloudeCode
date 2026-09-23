"""Gera a relação dos processos da 9ª DPE Ribeirão Preto que tramitam no eproc do TJSP.

Uso: python gerar_relacao_eproc.py [caminho_saida.xlsx]
Usa a base comum (base_dol.py): PAs ativos da 9ª no DOL, inclusive os que aparecem só como correlatos dentro de
outro PA, com a marcação de sistema do DOL e as certidões de migração lidas nas movimentações do e-SAJ.
"""
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from base_dol import RELEITURA, carregar, dt, foro_vara, representados

SAIDA_PADRAO = Path.home() / "Downloads" / "3. Processos eproc 9ª DPE" / "Processos eproc 9ª DPE.xlsx"

VERDE = "009632"
FINO = Side(style="thin", color="D9D9D9")
BORDA = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)

INICIADO = "Iniciado no eproc"
MIGRADO_OK = "Migrado do e-SAJ (DOL já indica eproc)"
MIGRADO_ESAJ = "Migrado do e-SAJ (DOL ainda indica e-SAJ)"
INCIDENTE = "Incidente migrado com número próprio"
EPROC_SEM_REG = "eproc no DOL, sem certidão de migração"
CATEGORIAS = [INICIADO, MIGRADO_OK, MIGRADO_ESAJ, INCIDENTE, EPROC_SEM_REG]
CORES = {INICIADO: "1F4E79", MIGRADO_OK: "375623", MIGRADO_ESAJ: "C55A11", INCIDENTE: "7030A0", EPROC_SEM_REG: "7F6000"}
FUNDOS = {INICIADO: "DDEBF7", MIGRADO_OK: "E2EFDA", MIGRADO_ESAJ: "FCE4D6", INCIDENTE: "E4DFEC", EPROC_SEM_REG: "FFF2CC"}
ESAJ = "e-SAJ, sem certidão de migração"
ESAJ_INCIDENTE = "e-SAJ; incidente vinculado foi para o eproc"
ESAJ_SEM_MOV = "e-SAJ no DOL, sem movimentações para conferir"
FORA_TJSP = "Fora do TJSP"
SEM_CNJ = "Sem número de processo no DOL"


def migr_principal(p):
    return [m for m in p["migr"] if not m["numero"] or m["numero"] == p["cnj"]]


def migr_incidentes(p):
    return [m for m in p["migr"] if m["numero"] and m["numero"] != p["cnj"]]


def ultima(datas):
    return max(datas, key=dt) if datas else ""


def itens_eproc(p):
    """Números que estão no eproc a partir de um PA: o próprio número e os incidentes migrados com número próprio."""
    itens = []
    if not p["cnj"]:
        return itens
    principal = migr_principal(p)
    if p["sis"] == "eproc":
        cat = MIGRADO_OK if principal else INICIADO if p["sem_mov"] else EPROC_SEM_REG
        itens.append((p["cnj"], cat, ultima([m["data"] for m in principal])))
    elif principal:
        itens.append((p["cnj"], MIGRADO_ESAJ, ultima([m["data"] for m in principal])))
    for m in migr_incidentes(p):
        itens.append((m["numero"], INCIDENTE, m["data"]))
    return itens


def resultado(p):
    if not p["cnj"]:
        return SEM_CNJ
    proprio = [cat for num, cat, _ in itens_eproc(p) if num == p["cnj"]]
    if proprio:
        return proprio[0]
    if migr_incidentes(p):
        return ESAJ_INCIDENTE
    if ".8.26." not in p["cnj"]:
        return FORA_TJSP
    return ESAJ_SEM_MOV if p["sem_mov"] else ESAJ


def relacao(pas):
    grupos = {}
    for p in pas:
        for numero, cat, data in itens_eproc(p):
            g = grupos.setdefault(numero, {"numero": numero, "itens": []})
            g["itens"].append({"pa": p, "cat": cat, "data": data})
    for g in grupos.values():
        proprias = [i["cat"] for i in g["itens"] if i["cat"] != INCIDENTE]
        if MIGRADO_ESAJ in proprias:
            g["cat"] = MIGRADO_ESAJ
        else:
            g["cat"] = proprias[0] if proprias else INCIDENTE
        g["data"] = ultima([i["data"] for i in g["itens"] if i["data"]])
    return grupos


def celula(ws, r, c, v, bold=False, cor=None, fundo=None, h="left", size=10, wrap=True, borda=True, v_al="top"):
    x = ws.cell(row=r, column=c, value=v)
    x.font = Font(bold=bold, color=cor, size=size, name="Calibri")
    if fundo:
        x.fill = PatternFill("solid", fgColor=fundo)
    x.alignment = Alignment(horizontal=h, vertical=v_al, wrap_text=wrap)
    if borda:
        x.border = BORDA
    return x


def faixa(ws, titulo, subtitulo, ultima_col):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ultima_col)
    celula(ws, 1, 1, titulo, bold=True, cor="FFFFFF", fundo=VERDE, size=15, borda=False, v_al="center")
    ws.row_dimensions[1].height = 32
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ultima_col)
    celula(ws, 2, 1, subtitulo, cor="595959", size=10, borda=False, wrap=False)
    ws.sheet_view.showGridLines = False


def cabecalho(ws, linha, colunas):
    for i, (h, w) in enumerate(colunas, 1):
        celula(ws, linha, i, h, bold=True, cor="FFFFFF", fundo=VERDE, h="center", v_al="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[linha].height = 44


def colorir(serie, cores):
    serie.graphicalProperties.line.noFill = True
    for idx, cor in enumerate(cores):
        ponto = DataPoint(idx=idx)
        ponto.graphicalProperties.solidFill = cor
        ponto.graphicalProperties.line.noFill = True
        serie.dPt.append(ponto)


def rotulos_barras(serie):
    serie.dLbls = DataLabelList()
    serie.dLbls.showVal = True
    for k in ("showSerName", "showCatName", "showLegendKey", "showPercent"):
        setattr(serie.dLbls, k, False)


def observacoes(g):
    obs = []
    marcas = {i["pa"]["sis"] for i in g["itens"] if i["cat"] != INCIDENTE}
    if len(marcas) > 1:
        obs.append("O número tem mais de um PA na 9ª e o DOL marca o sistema de forma diferente em cada um: " + "; ".join(
            f"PA {i['pa']['pa'] or i['pa']['id']} marcado como {i['pa']['sis']}" for i in g["itens"]
            if i["cat"] != INCIDENTE))
    for i in g["itens"]:
        p = i["pa"]
        if i["cat"] == INCIDENTE and g["cat"] == INCIDENTE:
            if migr_principal(p):
                situacao = "; o número cadastrado no PA também foi para o eproc e tem linha própria"
            elif p["ult"]:
                situacao = (f"; o número cadastrado no PA continua no e-SAJ, com andamento de maior sequência "
                            f"'{p['ult']}'")
            else:
                situacao = "; o número cadastrado no PA continua no e-SAJ"
            obs.append(f"A certidão de migração deste número está nas movimentações do PA {p['pa'] or p['id']}, "
                       f"cadastrado com o número {p['cnj']}{situacao}")
        elif i["cat"] == INCIDENTE:
            obs.append(f"Também é citado como incidente migrado nas movimentações do PA {p['pa'] or p['id']} "
                       f"(número {p['cnj']})")
            continue
        if i["cat"] == MIGRADO_ESAJ:
            obs.append(f"No PA {p['pa'] or p['id']} o DOL ainda mostra o número como e-SAJ")
        datas = sorted({m["data"] for m in migr_principal(p)}, key=dt) if i["cat"] != INCIDENTE else []
        if len(datas) > 1:
            obs.append(f"O e-SAJ registra mais de uma certidão de migração ({', '.join(datas)})")
        if p["pai"]:
            obs.append(f"PA {p['pa']} aparece como correlato dentro da linha de outro PA ({p['origem']})")
        if p["novo"]:
            obs.append(f"PA novo na 9ª, distribuído em {p['novo']}")
        if p["redist"]:
            obs.append(f"PA redistribuído à 9ª em {p['redist'][0]} (origem: {p['redist'][1]})")
    if g["cat"] == INICIADO:
        obs.append("Sem nenhuma movimentação do e-SAJ no DOL: o processo já nasceu no eproc")
    return ". ".join(dict.fromkeys(obs))


def aba_relacao(wb, grupos):
    ws = wb.create_sheet("Processos eproc")
    cols = [("Nº", 5), ("Número do processo no eproc", 27), ("Nº do PA no DOL", 14),
            ("Número cadastrado no PA (se diferente)", 25), ("Origem no eproc", 26), ("Remessa ao eproc", 12),
            ("Sistema indicado no DOL", 12), ("Foro / Vara (DOL)", 34), ("Ação (DOL)", 24),
            ("Representados pela Defensoria (DOL)", 38), ("Pasta Digital Eproc no DOL", 22), ("Observações", 62)]
    faixa(ws, "Processos no eproc  |  9ª Defensoria da Unidade Ribeirão Preto",
          f"{len(grupos)} números de processo   |   Um número por linha   |   Use os filtros do cabeçalho", len(cols))
    cabecalho(ws, 4, cols)
    ordem = {c: i for i, c in enumerate(CATEGORIAS)}
    lista = sorted(grupos.values(), key=lambda g: (ordem[g["cat"]], g["numero"]))
    for n, g in enumerate(lista, 1):
        r = 4 + n
        proprios = [i for i in g["itens"] if i["cat"] != INCIDENTE] or g["itens"]
        p = proprios[0]["pa"]
        pas = list(dict.fromkeys(i["pa"]["pa"] or i["pa"]["id"] for i in proprios))
        cadastrados = [i["pa"]["cnj"] for i in proprios if i["pa"]["cnj"] != g["numero"]]
        if g["cat"] == INCIDENTE:
            sistema, pasta = "Sem PA próprio", "Sem PA próprio no DOL"
        else:
            marcas = list(dict.fromkeys(i["pa"]["sis"] for i in proprios))
            sistema = " e ".join(marcas)
            if all(m == "eproc" for m in marcas):
                pasta = ("Aviso de acesso indisponível" if any(i["pa"]["sem_acesso"] for i in proprios)
                         else "Não verificado na lista (PA correlato)")
            elif "eproc" in marcas:
                pasta = "Varia entre os PAs (ver observações)"
            else:
                pasta = "Sem botão (DOL trata como e-SAJ)"
        vals = [n, g["numero"], "\n".join(pas), "\n".join(dict.fromkeys(cadastrados)) or "-", g["cat"], g["data"] or "-",
                sistema, foro_vara(p), p["acao"],
                "\n".join(dict.fromkeys(representados(i["pa"]) for i in proprios)), pasta, observacoes(g)]
        fundo = "F7F7F7" if n % 2 == 0 else None
        for c, v in enumerate(vals, 1):
            celula(ws, r, c, v, fundo=fundo, bold=c == 2, h="center" if c in (1, 3, 6, 7) else "left")
        celula(ws, r, 5, g["cat"], bold=True, cor=CORES[g["cat"]], fundo=FUNDOS[g["cat"]])
        if pasta.startswith("Aviso"):
            ws.cell(row=r, column=11).font = Font(color="C00000", size=10, name="Calibri")
        texto = max(len(vals[11]) / 58, len(vals[9]) / 36, len(vals[7]) / 32)
        ws.row_dimensions[r].height = min(max(30, 14 * (int(texto) + 1)), 200)
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{4 + len(lista)}"
    return lista


def aba_resumo(wb, meta, grupos, pas):
    ws = wb.active
    ws.title = "Resumo"
    faixa(ws, "Processos no eproc  |  9ª Defensoria Ribeirão Preto",
          f"Leitura do DOL de {meta['extraido_em']}   |   Somente leitura: nada foi alterado no DOL   |   "
          f"Consulta feita com o usuário {meta['usuario']}", 10)
    for col, w in zip("ABCDEFGHIJ", [40, 12, 10, 14, 14, 14, 14, 14, 14, 14]):
        ws.column_dimensions[col].width = w

    cont = Counter(g["cat"] for g in grupos.values())
    migrados = cont[MIGRADO_OK] + cont[MIGRADO_ESAJ] + cont[INCIDENTE]
    lista = [p for p in pas if not p["pai"]]
    cnjs = {p["cnj"] for p in pas if p["cnj"]}
    no_eproc = {p["cnj"] for p in pas if resultado(p) in CATEGORIAS}
    cards = [("Processos no eproc", len(grupos), "FFFFFF", VERDE, 1),
             ("Iniciados no eproc", cont[INICIADO], "FFFFFF", CORES[INICIADO], 2),
             ("Migrados do e-SAJ", migrados, "FFFFFF", "375623", 4),
             ("Migrados que o DOL\nainda mostra como e-SAJ", cont[MIGRADO_ESAJ], "FFFFFF", CORES[MIGRADO_ESAJ], 6),
             ("Incidentes com número\npróprio no eproc", cont[INCIDENTE], "FFFFFF", CORES[INCIDENTE], 8)]
    for rot, val, fg, bg, c in cards:
        if c > 1:
            ws.merge_cells(start_row=4, start_column=c, end_row=4, end_column=c + 1)
            ws.merge_cells(start_row=5, start_column=c, end_row=5, end_column=c + 1)
        celula(ws, 4, c, rot, bold=True, cor=fg, fundo=bg, h="center", size=9, v_al="center")
        celula(ws, 5, c, val, bold=True, cor=fg, fundo=bg, h="center", size=20, v_al="center")
    ws.row_dimensions[4].height = 30
    ws.row_dimensions[5].height = 40
    celula(ws, 6, 1, f"Base: {len(pas)} PAs ativos da 9ª ({len(lista)} na lista de processos ativos e "
           f"{len(pas) - len(lista)} que aparecem só como correlatos dentro de outro PA), com {len(cnjs)} números de "
           f"processo distintos. Deles, {len(no_eproc)} ({len(no_eproc) / len(cnjs):.0%}) já estão no eproc; somam-se "
           f"{cont[INCIDENTE]} incidentes que foram para o eproc com número próprio.",
           cor="404040", size=10, borda=False, wrap=False)

    celula(ws, 8, 1, "Quantidade de processos no eproc, por origem", bold=True, cor=VERDE, size=12, borda=False,
           wrap=False)
    for i, h in enumerate(["Origem", "Processos", "%"], 1):
        celula(ws, 9, i, h, bold=True, cor="FFFFFF", fundo=VERDE, h="center", v_al="center")
    cats = [c for c in CATEGORIAS if cont[c] or c != EPROC_SEM_REG]
    r = 10
    for cat in cats:
        celula(ws, r, 1, cat, bold=True, cor=CORES[cat], fundo=FUNDOS[cat])
        celula(ws, r, 2, cont[cat], h="center", bold=True)
        celula(ws, r, 3, cont[cat] / len(grupos), h="center").number_format = "0%"
        r += 1
    celula(ws, r, 1, "Total", bold=True, fundo="F2F2F2")
    celula(ws, r, 2, len(grupos), bold=True, fundo="F2F2F2", h="center")
    celula(ws, r, 3, 1, bold=True, fundo="F2F2F2", h="center").number_format = "0%"
    fim_cat = r

    graf = BarChart()
    graf.type = "bar"
    graf.style = 10
    graf.title = "Processos da 9ª no eproc, por origem"
    graf.add_data(Reference(ws, min_col=2, min_row=9, max_row=fim_cat - 1), titles_from_data=True)
    graf.set_categories(Reference(ws, min_col=1, min_row=10, max_row=fim_cat - 1))
    graf.legend = None
    graf.x_axis.scaling.orientation = "maxMin"
    graf.x_axis.delete = False
    graf.y_axis.delete = True
    graf.y_axis.majorGridlines = None
    graf.gapWidth = 50
    colorir(graf.series[0], [CORES[c] for c in cats])
    rotulos_barras(graf.series[0])
    graf.height, graf.width = 7.5, 17
    ws.add_chart(graf, "E8")

    r = fim_cat + 4
    celula(ws, r, 1, "Acervo ativo da 9ª por sistema (números de processo distintos)", bold=True, cor=VERDE, size=12,
           borda=False, wrap=False)
    r += 1
    ini_sis = r
    for i, h in enumerate(["Sistema", "Processos", "%"], 1):
        celula(ws, r, i, h, bold=True, cor="FFFFFF", fundo=VERDE, h="center", v_al="center")
    por_cnj = {}
    for p in pas:
        if p["cnj"]:
            res = resultado(p)
            atual = por_cnj.get(p["cnj"])
            por_cnj[p["cnj"]] = res if atual is None or res in CATEGORIAS else atual
    sistemas = Counter("eproc" if c in CATEGORIAS else FORA_TJSP if c == FORA_TJSP else "e-SAJ"
                       for c in por_cnj.values())
    linhas = [("eproc", sistemas["eproc"]), ("e-SAJ", sistemas["e-SAJ"]), (FORA_TJSP, sistemas[FORA_TJSP])]
    for rot, v in linhas:
        r += 1
        celula(ws, r, 1, rot, bold=True)
        celula(ws, r, 2, v, h="center", bold=True)
        celula(ws, r, 3, v / len(por_cnj), h="center").number_format = "0%"
    r += 1
    celula(ws, r, 1, "Total de números", bold=True, fundo="F2F2F2")
    celula(ws, r, 2, len(por_cnj), bold=True, fundo="F2F2F2", h="center")
    celula(ws, r, 3, 1, bold=True, fundo="F2F2F2", h="center").number_format = "0%"
    sem_cnj = sum(1 for p in pas if not p["cnj"])
    if sem_cnj:
        celula(ws, r + 1, 1, f"Além deles, {sem_cnj} PA ativo sem número de processo no DOL.", cor="808080", size=9,
               borda=False, wrap=False)

    graf2 = BarChart()
    graf2.type = "col"
    graf2.style = 10
    graf2.title = "Acervo ativo da 9ª por sistema"
    graf2.add_data(Reference(ws, min_col=2, min_row=ini_sis, max_row=ini_sis + len(linhas)), titles_from_data=True)
    graf2.set_categories(Reference(ws, min_col=1, min_row=ini_sis + 1, max_row=ini_sis + len(linhas)))
    graf2.legend = None
    graf2.y_axis.delete = True
    graf2.y_axis.majorGridlines = None
    graf2.x_axis.delete = False
    graf2.gapWidth = 80
    colorir(graf2.series[0], [VERDE, "7F7F7F", "BFBFBF"])
    rotulos_barras(graf2.series[0])
    graf2.height, graf2.width = 7.5, 17
    ws.add_chart(graf2, f"E{fim_cat + 9}")

    r = max(r + 3, fim_cat + 26)
    celula(ws, r, 1, "Remessas do e-SAJ ao eproc por mês", bold=True, cor=VERDE, size=12, borda=False, wrap=False)
    r += 1
    for i, h in enumerate(["Mês", "Processos"], 1):
        celula(ws, r, i, h, bold=True, cor="FFFFFF", fundo=VERDE, h="center", v_al="center")
    meses = Counter(g["data"][3:] for g in grupos.values() if g["data"])
    for mes in sorted(meses, key=lambda m: (m[3:], m[:2])):
        r += 1
        celula(ws, r, 1, mes, h="center")
        celula(ws, r, 2, meses[mes], h="center", bold=True)
    ws.freeze_panes = "A4"
    return cont, sistemas


def primeira_remessa(pas):
    return min((dt(m["data"]) for p in pas for m in p["migr"] if m["data"]), default=None)


def andamento_antigo(p, limite):
    return bool(limite and resultado(p) == ESAJ and p["ult_data"] and dt(p["ult_data"]) < limite)


def aba_conferencia(wb, pas):
    ws = wb.create_sheet("Conferência (todos)")
    limite = primeira_remessa(pas)
    cols = [("Número cadastrado no PA", 26), ("Nº do PA", 14), ("Onde aparece no DOL", 30),
            ("Sistema indicado no DOL", 12), ("Resultado da conferência", 38), ("Número no eproc (se diferente)", 25),
            ("Remessa ao eproc", 12), ("Andamento mais recente no e-SAJ (data)", 13), ("Foro / Vara (DOL)", 40),
            ("Ação (DOL)", 36), ("Observação", 40)]
    faixa(ws, "Conferência de todos os PAs ativos da 9ª",
          f"{len(pas)} PAs   |   Cada linha é um PA do DOL   |   Filtre a coluna E para ver cada grupo", len(cols))
    cabecalho(ws, 4, cols)
    ordem = {c: i for i, c in enumerate(CATEGORIAS + [ESAJ_INCIDENTE, ESAJ, ESAJ_SEM_MOV, FORA_TJSP, SEM_CNJ])}
    lista = sorted(pas, key=lambda p: (ordem[resultado(p)], p["cnj"] or "~", p["id"]))
    for n, p in enumerate(lista, 1):
        r = 4 + n
        res = resultado(p)
        itens = itens_eproc(p)
        outros = [num for num, cat, _ in itens if num != p["cnj"]]
        remessa = ultima([d for _, _, d in itens if d])
        obs = []
        if p["novo"]:
            obs.append(f"PA novo na 9ª, distribuído em {p['novo']}")
        if p["redist"]:
            obs.append(f"Redistribuído à 9ª em {p['redist'][0]} (origem: {p['redist'][1]})")
        antigo = andamento_antigo(p, limite)
        if antigo:
            obs.append(f"Andamento mais recente no e-SAJ anterior à primeira remessa ao eproc encontrada "
                       f"({limite:%d/%m/%Y}): vale conferir no tribunal")
        vals = [p["cnj"] or "(sem número)", p["pa"] or p["id"], p["origem"],
                {"eproc": "eproc", "e-SAJ": "e-SAJ"}.get(p["sis"], "sem indicação"), res, "\n".join(outros) or "-",
                remessa or "-", "-" if p["sem_mov"] else p["ult_data"] or "-", foro_vara(p), p["acao"],
                ". ".join(obs)]
        fundo = "F7F7F7" if n % 2 == 0 else None
        for c, v in enumerate(vals, 1):
            celula(ws, r, c, v, fundo=fundo, h="center" if c in (2, 4, 7, 8) else "left", wrap=c in (3, 6, 9, 10, 11))
        if res in FUNDOS:
            celula(ws, r, 5, res, bold=True, cor=CORES[res], fundo=FUNDOS[res])
        elif res == ESAJ_INCIDENTE:
            celula(ws, r, 5, res, bold=True, cor=CORES[INCIDENTE], fundo=fundo)
        if antigo:
            ws.cell(row=r, column=11).font = Font(color="C00000", size=10, name="Calibri")
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{4 + len(lista)}"


def aba_arquivados(wb, arquivados):
    ws = wb.create_sheet("Arquivados no DOL (eproc)")
    cols = [("Número do processo", 26), ("Nº do PA", 14), ("PA arquivado em", 13), ("Remessa ao eproc", 13),
            ("Foro / Vara (DOL)", 40), ("Ação (DOL)", 26), ("Representados pela Defensoria (DOL)", 40)]
    faixa(ws, "PAs já arquivados na 9ª que o DOL indica como eproc",
          "Apenas para informação: não entram na relação principal", len(cols))
    cabecalho(ws, 4, cols)
    for n, p in enumerate(arquivados, 1):
        vals = [p["cnj"], p["pa"] or p["id"], p["data_arq"], ultima([m["data"] for m in p["migr"]]) or "-",
                foro_vara(p), p["acao"], representados(p)]
        for c, v in enumerate(vals, 1):
            celula(ws, 4 + n, c, v, h="center" if c in (2, 3, 4) else "left")
    if not arquivados:
        celula(ws, 5, 1, "Nenhum.", cor="808080", borda=False)


def aba_metodo(wb, meta, pas):
    limite = primeira_remessa(pas)
    antigos = sum(1 for p in pas if andamento_antigo(p, limite))
    com_cert = sum(1 for p in pas if not p["pai"] and p["migr"])
    n_cert = sum(len(p["migr"]) for p in pas if not p["pai"])
    t = meta["totais"]
    ws = wb.create_sheet("Como foi feito")
    faixa(ws, "Como a relação foi montada", "Consulta feita somente em modo leitura no DOL", 1)
    ws.column_dimensions["A"].width = 135
    linhas = [
        f"Fonte: DOL, Acompanhamento de Processo da 9ª Defensoria da UNIDADE RIBEIRÃO PRETO. Primeira leitura em "
        f"{meta['extraido_primeira']} e leitura completa de conferência em {meta['extraido_em']}; esta planilha usa a "
        f"segunda, com releitura complementar da lista de ativos em {RELEITURA}, para o número dos PAs que têm correlatos. Lista de processos ativos: {t['ATIVO']} PAs (os {t['NOVO']} novos e os {t['REDISTRIBUIDO']} "
        f"redistribuídos já estão entre eles). Na mesma lista, {t['CORRELATOS']} PAs aparecem apenas na coluna "
        f"'Processos correlatos' de outro PA; eles também foram lidos um a um e entram na base, que soma {len(pas)} PAs.",
        "Sistema indicado no DOL: na lista, o DOL mostra o botão 'Pasta Digital Eproc' para os números que reconhece "
        "como eproc; na ficha do PA, o campo do número traz a legenda EPROC ou SAJ. Para os PAs da lista foi usado o "
        "botão; para os correlatos, a legenda da ficha. Numa amostra de 6 PAs as duas marcações coincidiram.",
        "Processos que saíram do e-SAJ: para cada número foram lidas as movimentações do e-SAJ guardadas no DOL. A "
        "certidão 'Remetidos os autos em razão de migração para outro sistema' informa que o processo passa a tramitar "
        "no eproc do TJSP e traz o número. Todas as certidões encontradas mencionam o eproc.",
        "Conferência independente: na segunda leitura, as certidões foram procuradas no texto bruto da página, sem "
        f"depender da tabela. O resultado dos PAs da lista foi idêntico ao da primeira leitura ({com_cert} PAs com "
        f"certidão, {n_cert} certidões, os mesmos números). Entre as duas leituras, o DOL passou a marcar como eproc o PA 6334379/2026 "
        "(0021253-10.2025.8.26.0506), que já constava como migrado.",
        "Incidentes: a tela de movimentações do DOL junta os andamentos do processo principal com os dos incidentes. "
        "Quando a certidão de migração traz número diferente do cadastrado no PA, quem foi para o eproc foi o incidente, "
        "com esse número próprio.",
        f"{INICIADO}: o DOL indica eproc e não há nenhuma movimentação do e-SAJ para o número.",
        f"{MIGRADO_OK}: há a certidão de migração do próprio número e o DOL já indica eproc.",
        f"{MIGRADO_ESAJ}: há a certidão de migração do próprio número, mas o DOL ainda indica e-SAJ em pelo menos um "
        "dos PAs desse número.",
        f"{INCIDENTE}: incidente de um processo acompanhado pela 9ª que foi para o eproc com número diferente do "
        "cadastrado no PA.",
        "Menções ao eproc sem certidão de migração (carta precatória para tribunal que só recebe petição pelo eproc) "
        "foram conferidas uma a uma e não indicam mudança de sistema.",
        "Pasta Digital Eproc no DOL: nos PAs da lista que o DOL reconhece como eproc, aparece ao usuário da consulta o "
        "aviso 'O seu acesso à pasta digital deste processo não está disponível'. O aviso depende do perfil de quem "
        "consulta. Para PAs correlatos essa informação não aparece na lista.",
        "Foro / Vara: copiados do cadastro do PA no DOL; quando o DOL não informa, a célula diz isso. Representados pela "
        "Defensoria: nomes da coluna Nome da lista de ativos; para os PAs sem nome na lista e para os correlatos, as "
        "partes que a ficha do PA marca como 'Representado por Defensoria'. Nenhum CPF foi copiado.",
        f"Pontos para conferir: {antigos} processos classificados como e-SAJ têm o andamento mais recente no DOL "
        f"anterior à primeira remessa ao eproc encontrada ({limite:%d/%m/%Y}); estão marcados em vermelho na aba "
        "'Conferência (todos)'. O DOL não guarda movimentações do eproc, então a data da remessa é a da certidão no "
        "e-SAJ. PAs arquivados não tiveram as movimentações conferidas; os que o DOL marca como eproc estão em aba "
        "própria. Processos de outros tribunais não têm movimentações no DOL.",
        "Nenhum dado foi incluído, alterado ou excluído no DOL: foram feitas apenas consultas de leitura.",
    ]
    for i, texto in enumerate(linhas, 4):
        celula(ws, i, 1, f"{i - 3}. {texto}", borda=False, size=11)
        ws.row_dimensions[i].height = 16 * (len(texto) // 120 + 1) + 6


def main():
    saida = Path(sys.argv[1]) if len(sys.argv) > 1 else SAIDA_PADRAO
    saida.parent.mkdir(parents=True, exist_ok=True)
    meta, pas, arquivados = carregar()
    grupos = relacao(pas)

    wb = Workbook()
    cont, sistemas = aba_resumo(wb, meta, grupos, pas)
    aba_relacao(wb, grupos)
    aba_conferencia(wb, pas)
    aba_arquivados(wb, arquivados)
    aba_metodo(wb, meta, pas)

    try:
        wb.save(saida)
        print(f"OK: {saida}")
    except PermissionError:
        alt = saida.with_name(f"{saida.stem} ({datetime.now().strftime('%Y-%m-%d %Hh%M')}){saida.suffix}")
        wb.save(alt)
        print(f"Planilha principal aberta no Excel; salvo em: {alt}")
    print(f"Números no eproc: {len(grupos)} | {dict(cont)}")
    print(f"Acervo por sistema: {dict(sistemas)}")
    print(f"Resultado por PA: {dict(Counter(resultado(p) for p in pas))}")


if __name__ == "__main__":
    main()

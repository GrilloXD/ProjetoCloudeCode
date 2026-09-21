"""Gera a planilha de controle de prazos das providências do DOL.

Uso: python gerar_planilha.py [caminho_saida.xlsx]
Lê dados/providencias.json e dados/feriados.json; acrescenta a execução em dados/historico.json.
"""
import json
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = Path(__file__).parent
DADOS = BASE / "dados"
SAIDA_PADRAO = Path(r"C:\Users\isabe\OneDrive\Área de Trabalho\Controle de Prazos DOL.xlsx")

VERDE = "009632"
CORES = {
    "VENCIDO": ("C00000", "FFFFFF"),
    "CRÍTICO": ("F4B183", "000000"),
    "ATENÇÃO": ("FFE699", "000000"),
    "EM DIA": ("C6EFCE", "000000"),
    "SEM PRAZO": ("E7E6E6", "000000"),
}


def br(d):
    return d.strftime("%d/%m/%Y") if d else ""


def parse(s):
    return datetime.strptime(s, "%d/%m/%Y").date() if s else None


def carregar_feriados():
    cfg = json.loads((DADOS / "feriados.json").read_text(encoding="utf-8"))
    dias = {parse(d) for d in cfg["feriados"]}
    for ano in range(2025, 2029):
        ini, fim = date(ano, 12, 20), date(ano + 1, 1, 20)
        d = ini
        while d <= fim:
            dias.add(d)
            d += timedelta(days=1)
    return dias


FERIADOS = carregar_feriados()


def util(d):
    return d.weekday() < 5 and d not in FERIADOS


def proximo_util(d):
    d += timedelta(days=1)
    while not util(d):
        d += timedelta(days=1)
    return d


def somar_uteis(inicio_contagem, n):
    d = inicio_contagem
    while not util(d):
        d += timedelta(days=1)
    contados = 1
    while contados < n:
        d = proximo_util(d)
        contados += 1
    return d


def uteis_entre(a, b):
    if b < a:
        return -uteis_entre(b, a)
    n, d = 0, a
    while d < b:
        d += timedelta(days=1)
        if util(d):
            n += 1
    return n


def calcular(p, hoje):
    """Retorna (data_intimacao, vencimento, observacao_calculo)."""
    if not p.get("prazo_judicial_dias"):
        return None, None, "Sem prazo processual em curso"
    intim = parse(p.get("data_intimacao"))
    obs = []
    if not intim and p.get("data_remessa_portal"):
        intim = parse(p["data_remessa_portal"]) + timedelta(days=10)
        obs.append(f"Intimação tácita projetada (remessa {p['data_remessa_portal']} + 10 dias corridos)")
    if not intim:
        return None, None, "Sem data de intimação identificada"
    dias = p["prazo_judicial_dias"] * (2 if p.get("prazo_em_dobro", True) else 1)
    venc = somar_uteis(proximo_util(intim), dias)
    obs.append(f"{p['prazo_judicial_dias']} dias{' em dobro' if p.get('prazo_em_dobro', True) else ''} = {dias} dias úteis")
    return intim, venc, "; ".join(obs)


def status(venc, hoje):
    if not venc:
        return "SEM PRAZO", None
    r = uteis_entre(hoje, venc)
    if venc < hoje:
        return "VENCIDO", r
    if r <= 5:
        return "CRÍTICO", r
    if r <= 10:
        return "ATENÇÃO", r
    return "EM DIA", r


def main():
    saida = Path(sys.argv[1]) if len(sys.argv) > 1 else SAIDA_PADRAO
    hoje = date.today()
    dados = json.loads((DADOS / "providencias.json").read_text(encoding="utf-8"))
    provs = dados["providencias"]

    linhas = []
    for p in provs:
        intim, venc, obs_calc = calcular(p, hoje)
        st, rest = status(venc, hoje)
        prov = parse(p.get("prazo_providencia"))
        if venc and prov:
            dif = uteis_entre(prov, venc)
            comp = "Providência antes do prazo real" if prov <= venc else "ATENÇÃO: providência DEPOIS do prazo real"
        elif venc and not prov:
            dif, comp = None, "Providência sem data: lançar o prazo real"
        else:
            dif, comp = None, ""
        linhas.append(dict(p=p, intim=intim, venc=venc, st=st, rest=rest, dif=dif, comp=comp, obs_calc=obs_calc))

    ordem = {"VENCIDO": 0, "CRÍTICO": 1, "ATENÇÃO": 2, "EM DIA": 3, "SEM PRAZO": 4}
    linhas.sort(key=lambda l: (ordem[l["st"]], l["venc"] or date.max))

    wb = Workbook()
    ws = wb.active
    ws.title = "Painel"
    fino = Side(style="thin", color="BFBFBF")
    borda = Border(left=fino, right=fino, top=fino, bottom=fino)

    ws.merge_cells("A1:R1")
    ws["A1"] = f"Controle de Prazos DOL | 9ª Defensoria RP | Responsável: {dados.get('responsavel', '')}"
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF", name="Calibri")
    ws["A1"].fill = PatternFill("solid", fgColor=VERDE)
    ws.merge_cells("A2:R2")
    ws["A2"] = (f"Atualizado em {datetime.now().strftime('%d/%m/%Y %H:%M')} | "
                f"{len(provs)} providências | Leitura apenas: nada é registrado, cumprido ou excluído no DOL")
    ws["A2"].font = Font(italic=True, size=10)

    resumo = {k: sum(1 for l in linhas if l["st"] == k) for k in ordem}
    col = 1
    for k, v in resumo.items():
        c = ws.cell(row=3, column=col, value=f"{k}: {v}")
        bg, fg = CORES[k]
        c.fill = PatternFill("solid", fgColor=bg)
        c.font = Font(bold=True, color=fg)
        c.alignment = Alignment(horizontal="center")
        ws.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col + 1)
        col += 2

    cab = ["Status", "Dias úteis restantes", "Vencimento REAL", "Prazo na PROVIDÊNCIA", "Comparação",
           "Assistido(s)", "Processo (CNJ)", "PA", "Vara", "Tipo", "Inserida por / em",
           "O que foi pedido", "Intimação da DPE", "Prazo judicial", "Cálculo",
           "Informações processuais", "Alertas", "Duplicada?"]
    larg = [12, 10, 13, 13, 24, 28, 26, 13, 22, 14, 22, 60, 16, 14, 30, 55, 45, 18]
    for i, (h, w) in enumerate(zip(cab, larg), 1):
        c = ws.cell(row=5, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=VERDE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = borda
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[5].height = 32

    for r, l in enumerate(linhas, 6):
        p = l["p"]
        intim_txt = br(l["intim"])
        if l["intim"] and not p.get("data_intimacao"):
            intim_txt += " (tácita projetada)"
        vals = [
            l["st"], l["rest"], br(l["venc"]), p.get("prazo_providencia") or "sem data", l["comp"],
            p.get("assistidos", ""), p.get("processo", ""), p.get("pa", ""), p.get("vara", ""), p.get("tipo", ""),
            f"{p.get('inserido_por', '')} em {p.get('data_insercao', '')}",
            p.get("pedido", ""), intim_txt, p.get("prazo_judicial_txt", ""), l["obs_calc"],
            p.get("info_processual", ""), p.get("alertas", ""), p.get("duplicada", ""),
        ]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = borda
            c.alignment = Alignment(vertical="top", wrap_text=True,
                                    horizontal="center" if i <= 4 or i in (8, 13) else "left")
        bg, fg = CORES[l["st"]]
        ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=bg)
        ws.cell(row=r, column=1).font = Font(bold=True, color=fg)
        if "DEPOIS" in l["comp"] or "sem data" in l["comp"]:
            ws.cell(row=r, column=5).font = Font(bold=True, color="C00000")
        if p.get("alertas"):
            ws.cell(row=r, column=17).font = Font(bold=True, color="C00000")
        if r % 2 == 0:
            for i in range(2, len(cab) + 1):
                ws.cell(row=r, column=i).fill = PatternFill("solid", fgColor="F2F8F4")

    ws.freeze_panes = "G6"
    ws.auto_filter.ref = f"A5:{get_column_letter(len(cab))}{5 + len(linhas)}"
    ws.sheet_view.zoomScale = 90

    hist_path = DADOS / "historico.json"
    hist = json.loads(hist_path.read_text(encoding="utf-8")) if hist_path.exists() else []
    hist.append({
        "execucao": datetime.now().strftime("%d/%m/%Y %H:%M"),
        "total": len(provs),
        **{k: v for k, v in resumo.items()},
        "mudancas": dados.get("mudancas_desde_ultima", ""),
    })
    hist_path.write_text(json.dumps(hist, ensure_ascii=False, indent=2), encoding="utf-8")

    wh = wb.create_sheet("Histórico")
    hcab = ["Execução", "Total", "Vencido", "Crítico", "Atenção", "Em dia", "Sem prazo", "Mudanças desde a última"]
    for i, h in enumerate(hcab, 1):
        c = wh.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=VERDE)
        wh.column_dimensions[get_column_letter(i)].width = 90 if i == 8 else 14
    for r, h in enumerate(reversed(hist), 2):
        for i, k in enumerate(["execucao", "total", "VENCIDO", "CRÍTICO", "ATENÇÃO", "EM DIA", "SEM PRAZO", "mudancas"], 1):
            c = wh.cell(row=r, column=i, value=h.get(k, ""))
            c.alignment = Alignment(vertical="top", wrap_text=True)
    wh.freeze_panes = "A2"

    wr = wb.create_sheet("Regras de cálculo")
    regras = [
        "Intimação pessoal da Defensoria pelo portal eletrônico: vale na data em que a intimação é aberta no DOL.",
        "Sem abertura, a intimação é tácita 10 dias corridos após a remessa ao portal (Lei 11.419/2006, art. 5º, §3º).",
        "O prazo começa no primeiro dia útil seguinte à intimação.",
        "Prazo em dobro para a Defensoria (CPC, art. 186) e contagem em dias úteis (CPC, art. 219).",
        "Recesso forense de 20/12 a 20/01 com prazos suspensos (CPC, art. 220).",
        "Feriados e suspensões considerados estão em scripts/prazos_dol/dados/feriados.json; conferir comunicados do TJSP.",
        "Quando o despacho não fixa prazo, o DOL pode trazer um número cadastrado; o prazo legal supletivo é de 5 dias (CPC, art. 218, §3º). Esses casos aparecem em Alertas.",
        "Duas intimações do mesmo ato: conta-se da primeira, por segurança.",
        "Status: VENCIDO (já passou), CRÍTICO (até 5 dias úteis), ATENÇÃO (6 a 10), EM DIA (mais de 10), SEM PRAZO (sem prazo processual).",
        "A planilha é só de leitura e conferência. Confirme sempre nos autos antes de protocolar.",
    ]
    wr.column_dimensions["A"].width = 130
    wr["A1"] = "Regras usadas no cálculo"
    wr["A1"].font = Font(bold=True, size=13, color=VERDE)
    for i, t in enumerate(regras, 3):
        wr.cell(row=i, column=1, value=f"{i - 2}. {t}").alignment = Alignment(wrap_text=True)

    try:
        wb.save(saida)
        print(f"OK: {saida}")
    except PermissionError:
        alt = saida.with_name(f"{saida.stem} ({datetime.now().strftime('%Y-%m-%d %Hh%M')}){saida.suffix}")
        wb.save(alt)
        print(f"Planilha principal aberta no Excel; salvo em: {alt}")
    for l in linhas:
        print(f"{l['st']:9} | {br(l['venc']) or '-':10} | prov {l['p'].get('prazo_providencia') or 'sem data':10} | {l['p']['assistidos'][:40]}")


if __name__ == "__main__":
    main()
